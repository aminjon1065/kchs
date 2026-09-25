"""DOCX отчёта через docxtpl (ADR-0022, ADR-0078).

Шаблон — документ A4 с колонтитулами, полями номера страницы и заголовком; docxtpl
подставляет в него название, параметры и колонтитулы (Jinja в колонтитулах и теле).
Тело дописывается в отрисованный документ по модели страницы печати
(`window.kchsPrint`): текст Tiptap — абзацами и списками, таблицы — таблицами,
графики и карты — картинками, снятыми со страницы печати. Значения уже посчитаны
страницей с правами получателя, движок только раскладывает их.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Mapping
from dataclasses import dataclass
from functools import lru_cache
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docxtpl import DocxTemplate

PAGE_WIDTH_MM = 210
PAGE_HEIGHT_MM = 297
# Размеры страниц, мм (ADR-0164): короткая и длинная сторона
PAGE_MM = {"A4": (PAGE_WIDTH_MM, PAGE_HEIGHT_MM), "A3": (297, 420)}
MARGIN_MM = 18
MONO_FONT = "Courier New"


@dataclass(frozen=True)
class DocxLabels:
    header: str
    footer: str
    page: str
    of: str


def _field(paragraph: Any, instruction: str) -> None:
    """Поле Word (номер страницы, число страниц): обновляется при открытии."""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


@lru_cache(maxsize=4)
def template_bytes(orientation: str, size: str = "A4") -> bytes:
    """Шаблон docxtpl: страница нужного размера и ориентации, колонтитулы, заголовок и тело."""
    doc = Document()
    section = doc.sections[0]
    landscape = orientation == "landscape"
    short, long = PAGE_MM.get(size, PAGE_MM["A4"])
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Mm(long if landscape else short)
    section.page_height = Mm(short if landscape else long)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(MARGIN_MM))

    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(10)

    header = section.header.paragraphs[0]
    header.add_run("{{ header }}").font.size = Pt(8)

    # Нижний колонтитул: текст слева, номер страницы — отдельной строкой у правого поля
    # (табуляция к правому краю в LibreOffice и Word ставится по-разному)
    footer = section.footer.paragraphs[0]
    footer.add_run("{{ footer }}").font.size = Pt(8)
    pages = section.footer.add_paragraph()
    pages.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pages.add_run("{{ page_label }} ").font.size = Pt(8)
    _field(pages, "PAGE")
    pages.add_run(" {{ of_label }} ").font.size = Pt(8)
    _field(pages, "NUMPAGES")

    doc.add_paragraph("{{ title }}", style="Title")
    doc.add_paragraph("{{ subtitle }}", style="Subtitle")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def content_width_mm(orientation: str, size: str = "A4") -> float:
    short, long = PAGE_MM.get(size, PAGE_MM["A4"])
    page = long if orientation == "landscape" else short
    return page - 2 * MARGIN_MM


# ─── Текст Tiptap → абзацы ──────────────────────────────────────────────────


def _children(node: Mapping[str, Any]) -> list[Mapping[str, Any]]:
    content = node.get("content")
    if not isinstance(content, list):
        return []
    return [child for child in content if isinstance(child, Mapping)]


def _add_inline(paragraph: Any, nodes: list[Mapping[str, Any]], *, code: bool = False) -> None:
    for node in nodes:
        kind = node.get("type")
        if kind == "hardBreak":
            paragraph.add_run().add_break(WD_BREAK.LINE)
            continue
        text = node.get("text")
        if kind != "text" or not isinstance(text, str):
            continue
        marks = {
            str(mark.get("type")): mark
            for mark in node.get("marks", []) or []
            if isinstance(mark, Mapping)
        }
        run = paragraph.add_run(text)
        run.bold = "bold" in marks or None
        run.italic = "italic" in marks or None
        run.underline = "underline" in marks or "link" in marks or None
        run.font.strike = "strike" in marks or None
        if code or "code" in marks:
            run.font.name = MONO_FONT
        link = marks.get("link")
        href = (link.get("attrs") or {}).get("href") if link else None
        if isinstance(href, str) and href.startswith(("http://", "https://", "mailto:")):
            paragraph.add_run(f" ({href})")


def _list_style(ordered: bool, depth: int) -> str:
    base = "List Number" if ordered else "List Bullet"
    return base if depth <= 1 else f"{base} {min(depth, 3)}"


def add_rich_text(target: Any, body: Mapping[str, Any]) -> None:
    """Документ Tiptap (белый список контракта `common/rich-text.ts`) — в документ Word."""

    def blocks(nodes: list[Mapping[str, Any]], depth: int = 0, style: str | None = None) -> None:
        for node in nodes:
            kind = node.get("type")
            if kind == "paragraph":
                paragraph = target.add_paragraph(style=style) if style else target.add_paragraph()
                _add_inline(paragraph, _children(node))
            elif kind == "heading":
                level = (node.get("attrs") or {}).get("level", 1)
                level = level if isinstance(level, int) and 1 <= level <= 3 else 1
                heading = target.add_heading(level=level)
                _add_inline(heading, _children(node))
            elif kind in ("bulletList", "orderedList"):
                ordered = kind == "orderedList"
                for item in _children(node):
                    item_style = _list_style(ordered, depth + 1)
                    blocks(_children(item), depth + 1, item_style)
            elif kind == "blockquote":
                blocks(_children(node), depth, "Quote")
            elif kind == "codeBlock":
                paragraph = target.add_paragraph()
                _add_inline(paragraph, _children(node), code=True)
            elif kind == "horizontalRule":
                target.add_paragraph("—" * 20).alignment = WD_ALIGN_PARAGRAPH.CENTER

    blocks(_children(body))


# ─── Модель страницы печати → тело DOCX ──────────────────────────────────────


def _caption(target: Any, text: str | None) -> None:
    if text:
        paragraph = target.add_paragraph()
        paragraph.add_run(text).bold = True


def _note(target: Any, text: str | None) -> None:
    if text:
        paragraph = target.add_paragraph()
        run = paragraph.add_run(text)
        run.italic = True
        run.font.size = Pt(8)


def _table(target: Any, block: Mapping[str, Any]) -> None:
    columns = [column for column in block.get("columns", []) if isinstance(column, Mapping)]
    rows = [row for row in block.get("rows", []) if isinstance(row, list)]
    if not columns:
        return
    table = target.add_table(rows=1 + len(rows), cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, column in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(column.get("label", "")))
        run.bold = True
        run.font.size = Pt(8)
    for row_index, row in enumerate(rows, start=1):
        cells = table.rows[row_index].cells
        for index, column in enumerate(columns):
            value = row[index] if index < len(row) else ""
            paragraph = cells[index].paragraphs[0]
            paragraph.add_run(str(value)).font.size = Pt(8)
            if column.get("numeric"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _metrics(target: Any, block: Mapping[str, Any]) -> None:
    items = [item for item in block.get("items", []) if isinstance(item, Mapping)]
    if not items:
        return
    table = target.add_table(rows=len(items), cols=3)
    table.style = "Table Grid"
    for index, item in enumerate(items):
        cells = table.rows[index].cells
        cells[0].text = str(item.get("label", ""))
        value = cells[1].paragraphs[0].add_run(str(item.get("value", "")))
        value.bold = True
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cells[2].text = str(item.get("note", ""))


ImageLookup = Callable[[str], bytes | None]


def build_docx(
    model: Mapping[str, Any],
    images: ImageLookup,
    labels: DocxLabels,
) -> bytes:
    """DOCX по модели страницы печати: картинки графиков и карт — из `images(id)`."""
    settings = model.get("settings") or {}
    orientation = "landscape" if settings.get("orientation") == "landscape" else "portrait"
    size = settings.get("pageSize") if settings.get("pageSize") in PAGE_MM else "A4"
    tpl = DocxTemplate(io.BytesIO(template_bytes(orientation, size)))
    tpl.render(
        {
            "title": str(model.get("title", "")),
            "subtitle": str(model.get("subtitle", "")),
            "header": labels.header,
            "footer": labels.footer,
            "page_label": labels.page,
            "of_label": labels.of,
        },
        autoescape=True,
    )
    # Тело — в отрисованный документ после заголовка
    body = tpl.docx
    width = content_width_mm(orientation, size)

    for block in model.get("blocks", []):
        if not isinstance(block, Mapping):
            continue
        kind = block.get("kind")
        if kind == "text":
            body_value = block.get("body")
            if isinstance(body_value, Mapping):
                add_rich_text(body, body_value)
        elif kind == "figure":
            _caption(body, block.get("title"))
            image = images(str(block.get("id")))
            if image:
                body.add_picture(io.BytesIO(image), width=Mm(width))
            _note(body, block.get("note"))
        elif kind == "table":
            _caption(body, block.get("title"))
            _table(body, block)
            total = block.get("total")
            shown = len(block.get("rows", []))
            if isinstance(total, int) and total > shown:
                _note(body, f"{shown} / {total}")
        elif kind == "metrics":
            _caption(body, block.get("title"))
            _metrics(body, block)
        elif kind == "notice":
            _caption(body, block.get("title"))
            _note(body, str(block.get("text", "")))
        elif kind == "page_break":
            body.add_page_break()

    out = io.BytesIO()
    tpl.save(out)
    return out.getvalue()
