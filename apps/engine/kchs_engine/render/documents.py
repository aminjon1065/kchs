"""Рендеры модуля документов (ADR-0085): задание `render:document.render`.

Движок не читает базу и не решает о правах: по идентификатору рендера он берёт
у api план (`/internal/documents/renders/<id>/start`), собранный с правами
заказчика в этот момент, и сообщает результат (`…/done`). Виды плана:

- `html` — печатная форма: страница, собранная api, печатается Chromium в PDF A4
  с колонтитулом «Лист N из M»; сеть страницы закрыта, скрипты выключены;
- `overlay` — штамп или водяной знак: страница наложения печатается под размер
  каждого листа (одна на каждый размер) и накладывается на PDF (pypdf); не-PDF
  сначала переводится в PDF, как PDF-представление версии. С картинками
  (факсимиле, ADR-0198) место каждой ищет движок — явное, по маркеру `[[SIGN]]`,
  левее фамилии на листах с конца, иначе запасное, — и наложение печатается для
  каждого такого листа отдельно; найденные места уходят в отчёт;
- `docx` — заполнение шаблона docxtpl в песочнице Jinja (`SandboxedEnvironment`,
  автоэкранирование XML): шаблон пишет администратор справочника, но из шаблона
  нельзя добраться до Python;
- `inspect` — разбор шаблона: пути плейсхолдеров (`doc.subject`,
  `doc.fields.addressee`) без переменных циклов.
"""

from __future__ import annotations

import asyncio
import html as html_lib
import io
import re
import subprocess
import tempfile
import time
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any

from kchs_engine.api import document_render_done, document_render_start
from kchs_engine.config import settings
from kchs_engine.contracts import document_render_contract
from kchs_engine.files.pdf import to_pdf
from kchs_engine.files.processing import ToolError
from kchs_engine.jobs.registry import handler
from kchs_engine.logging import log
from kchs_engine.render.report import (
    MARGIN_MM,
    count_pdf_pages,
    footer_template,
    render_slots,
    shared_browser,
)
from kchs_engine.storage import download, object_size, upload
from kchs_engine.tools import tool_env

if TYPE_CHECKING:
    from jinja2 import nodes as jinja_nodes

PDF_MIME = "application/pdf"


class RenderContentError(Exception):
    """Сбой из-за содержимого (шаблон, файл): повтор не поможет, рендер — «не удался».

    Текст — причина по-русски, её увидит пользователь.
    """


# ─── Разбор и заполнение шаблона DOCX ────────────────────────────────────────


def _sandbox() -> Any:
    """Песочница Jinja: без доступа к внутренностям Python; неизвестный путь
    (`{{ unknown.path }}`) — пустая строка, а не ошибка заполнения."""
    from jinja2 import ChainableUndefined
    from jinja2.sandbox import SandboxedEnvironment

    return SandboxedEnvironment(autoescape=True, undefined=ChainableUndefined)


def _unreadable() -> tuple[type[BaseException], ...]:
    """Ошибки чтения файла шаблона: не ZIP, не пакет Word, битый XML."""
    from docx.opc.exceptions import PackageNotFoundError
    from lxml.etree import XMLSyntaxError

    return (zipfile.BadZipFile, PackageNotFoundError, XMLSyntaxError, KeyError, ValueError)


def _template_source(path: Path) -> str:
    """Текст шаблона для Jinja: тело, верхние и нижние колонтитулы (как у docxtpl)."""
    from docx import Document
    from docx.oxml import parse_xml
    from docxtpl import DocxTemplate

    template = DocxTemplate(str(path))
    document = Document(str(path))
    xml = template.patch_xml(template.xml_to_string(document._element.body))
    for uri in (template.HEADER_URI, template.FOOTER_URI):
        for relation in document.part.rels.values():
            if relation.reltype == uri and relation.target_part.blob:
                part = template.xml_to_string(parse_xml(relation.target_part.blob))
                xml += template.patch_xml(part)
    return xml


def _chain(node: jinja_nodes.Node) -> list[str] | None:
    """`doc.fields.addressee` / `doc['fields']` → ['doc', 'fields', 'addressee']."""
    from jinja2 import nodes

    parts: list[str] = []
    current: Any = node
    while True:
        if isinstance(current, nodes.Getattr):
            parts.append(current.attr)
            current = current.node
        elif (
            isinstance(current, nodes.Getitem)
            and isinstance(current.arg, nodes.Const)
            and isinstance(current.arg.value, str)
        ):
            parts.append(current.arg.value)
            current = current.node
        elif isinstance(current, nodes.Name):
            parts.append(current.name)
            return list(reversed(parts))
        else:
            return None


def _names(target: Any) -> set[str]:
    """Имена, которые объявляет цикл или присваивание (`for a, b in …`)."""
    from jinja2 import nodes

    if isinstance(target, nodes.Name):
        return {target.name}
    if isinstance(target, nodes.Tuple):
        return set().union(*(_names(item) for item in target.items))
    return set()


def collect_placeholders(source: str) -> list[str]:
    """Пути плейсхолдеров шаблона: самые длинные цепочки атрибутов от корня контекста.

    Переменные циклов и присваиваний шаблона (и `loop`) — не контекст: их пути
    не возвращаются, а выражение цикла (`doc.attachments`) — возвращается.
    """
    from jinja2 import nodes

    found: set[str] = set()

    def visit(node: Any, bound: frozenset[str]) -> None:
        if isinstance(node, (nodes.Getattr, nodes.Getitem)):
            chain = _chain(node)
            if chain is not None:
                if chain[0] not in bound:
                    found.add(".".join(chain))
                return
        if isinstance(node, nodes.Name):
            if node.ctx == "load" and node.name not in bound:
                found.add(node.name)
            return
        if isinstance(node, nodes.For):
            visit(node.iter, bound)
            inner = bound | _names(node.target) | {"loop"}
            if node.test is not None:
                visit(node.test, inner)
            for child in [*node.body, *node.else_]:
                visit(child, inner)
            return
        if isinstance(node, nodes.Assign):
            visit(node.node, bound)
            return
        for child in node.iter_child_nodes():
            visit(child, bound)

    tree = _sandbox().parse(source)
    # Присваивания верхнего уровня видны дальше по шаблону
    bound: set[str] = set()
    for statement in tree.body:
        if isinstance(statement, nodes.Assign):
            bound |= _names(statement.target)
    visit(tree, frozenset(bound))
    return sorted(found)


def _template_error(error: Exception) -> RenderContentError:
    from jinja2 import TemplateError

    if isinstance(error, TemplateError):
        return RenderContentError(f"Ошибка в шаблоне: {error}")
    return RenderContentError("Файл шаблона не читается как документ Word (.docx)")


def inspect_template(path: Path) -> list[str]:
    """Разбор шаблона: пути плейсхолдеров; синтаксическая ошибка — причина для автора."""
    from jinja2 import TemplateError

    try:
        return collect_placeholders(_template_source(path))
    except (TemplateError, *_unreadable()) as error:
        raise _template_error(error) from error


def fill_template(path: Path, context: dict[str, Any], target: Path) -> Path:
    """Заполнение шаблона в песочнице: значения экранируются для XML Word."""
    from docxtpl import DocxTemplate
    from jinja2 import TemplateError

    try:
        template = DocxTemplate(str(path))
        template.render(context, jinja_env=_sandbox(), autoescape=True)
        template.save(str(target))
    except (TemplateError, *_unreadable()) as error:
        raise _template_error(error) from error
    return target


# ─── Наложение на PDF ────────────────────────────────────────────────────────


def _open_pdf(path: Path) -> Any:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise RenderContentError("PDF защищён паролем — наложение невозможно")
        _ = len(reader.pages)
    except PdfReadError as error:
        raise RenderContentError("PDF не читается") from error
    return reader


def target_pages(count: int, pages: str) -> list[int]:
    """Листы наложения: первый, последний или все."""
    if count == 0:
        return []
    if pages == "first":
        return [0]
    if pages == "last":
        return [count - 1]
    return list(range(count))


def page_sizes(path: Path, pages: str) -> list[tuple[float, float]]:
    """Размеры листов (pt) с учётом поворота: те, на которые ляжет наложение."""
    from pypdf import PdfWriter

    writer = PdfWriter(clone_from=_open_pdf(path))
    targets = target_pages(len(writer.pages), pages)
    sizes: list[tuple[float, float]] = []
    for index in targets:
        page = writer.pages[index]
        page.transfer_rotation_to_content()
        box = page.cropbox
        sizes.append((float(box.width), float(box.height)))
    return sizes


def size_key(size: tuple[float, float]) -> tuple[float, float]:
    return (round(size[0], 1), round(size[1], 1))


def merge_overlays(
    source: Path,
    overlays: dict[tuple[float, float], bytes],
    pages: str,
    target: Path,
) -> tuple[Path, int]:
    """Наложение на листы PDF: поворот листа переносится в содержимое, наложение
    масштабируется под лист (Chromium округляет размер до пикселя) и ставится в
    угол видимой области. Исходный файл не меняется — результат в `target`."""
    from pypdf import PdfReader, PdfWriter, Transformation

    writer = PdfWriter(clone_from=_open_pdf(source))
    targets = target_pages(len(writer.pages), pages)
    for index in targets:
        page = writer.pages[index]
        page.transfer_rotation_to_content()
        box = page.cropbox
        width, height = float(box.width), float(box.height)
        overlay = PdfReader(io.BytesIO(overlays[size_key((width, height))])).pages[0]
        scale_x = width / float(overlay.mediabox.width)
        scale_y = height / float(overlay.mediabox.height)
        page.merge_transformed_page(
            overlay,
            Transformation().scale(scale_x, scale_y).translate(float(box.left), float(box.bottom)),
        )
    with target.open("wb") as stream:
        writer.write(stream)
    return target, len(writer.pages)


# ─── Картинки наложения: где поставить факсимиле (ADR-0198) ───────────────────

MM = 72 / 25.4
DATA_URI = re.compile(r"^data:image/png;base64,[A-Za-z0-9+/]+={0,2}$")
WORD = re.compile(
    r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">(.*?)</word>'
)
PAGE = re.compile(r'<page width="([\d.]+)" height="([\d.]+)">(.*?)</page>', re.S)


@dataclass(frozen=True)
class PageBox:
    """Видимая область листа после переноса поворота (pt)."""

    left: float
    bottom: float
    width: float
    height: float


@dataclass(frozen=True)
class TextHit:
    """Найденный текст: лист и рамка в долях листа от левого верхнего угла."""

    page: int
    left: float
    top: float
    right: float
    bottom: float


@dataclass(frozen=True)
class Placement:
    """Центр картинки в долях листа и, у маркера, рамка под белую заливку."""

    page: int
    x: float
    y: float
    source: str
    cover: tuple[float, float, float, float] | None = None


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip().casefold()


def page_words(source: Path) -> list[list[tuple[str, float, float, float, float]]]:
    """Слова листов с рамками в долях листа — poppler `pdftotext -bbox`.

    pypdf не декодирует шрифты CID, которыми печатает Chromium; poppler читает их
    по ToUnicode и отдаёт рамки слов в видимой ориентации листа."""
    try:
        done = subprocess.run(
            ["pdftotext", "-bbox", "-enc", "UTF-8", str(source), "-"],
            capture_output=True,
            check=True,
            timeout=60,
            env=tool_env(),
        )
    except (OSError, subprocess.SubprocessError):
        log.warning("render.text_layer_unavailable")
        return []
    pages: list[list[tuple[str, float, float, float, float]]] = []
    for match in PAGE.finditer(done.stdout.decode("utf-8", "replace")):
        width, height = float(match.group(1)) or 1.0, float(match.group(2)) or 1.0
        words = [
            (
                html_lib.unescape(word.group(5)),
                float(word.group(1)) / width,
                float(word.group(2)) / height,
                float(word.group(3)) / width,
                float(word.group(4)) / height,
            )
            for word in WORD.finditer(match.group(3))
        ]
        pages.append(words)
    return pages


def find_text(
    pages: list[list[tuple[str, float, float, float, float]]], needle: str
) -> list[TextHit]:
    """Вхождения текста по словам строки: «Саидзода Т. Н.» — три слова одной строки."""
    wanted = _normalize(needle).split(" ")
    if not wanted or wanted == [""]:
        return []
    hits: list[TextHit] = []
    for index, words in enumerate(pages):
        lines: list[list[tuple[str, float, float, float, float]]] = []
        for word in sorted(words, key=lambda item: (item[2], item[1])):
            for line in lines:
                head = line[0]
                if (
                    abs((head[2] + head[4]) / 2 - (word[2] + word[4]) / 2)
                    <= (head[4] - head[2]) / 2
                ):
                    line.append(word)
                    break
            else:
                lines.append([word])
        for line in lines:
            line.sort(key=lambda item: item[1])
            texts = [_normalize(item[0]) for item in line]
            for start in range(len(line) - len(wanted) + 1):
                if texts[start : start + len(wanted)] != wanted:
                    continue
                matched = line[start : start + len(wanted)]
                hits.append(
                    TextHit(
                        index,
                        matched[0][1],
                        min(item[2] for item in matched),
                        matched[-1][3],
                        max(item[4] for item in matched),
                    )
                )
    return hits


def place_image(
    pages: list[list[tuple[str, float, float, float, float]]],
    boxes: list[PageBox],
    image: dict[str, Any],
) -> Placement:
    """Место картинки: выбранное человеком → маркер → текст на листах с конца → запасное."""
    last = len(boxes) - 1
    manual = image.get("manual")
    if manual:
        page = min(max(int(manual["page"]) - 1, 0), last)
        return Placement(page, float(manual["x"]), float(manual["y"]), "manual")
    for anchor in image.get("anchors") or []:
        hits = [hit for hit in find_text(pages, str(anchor["text"])) if hit.page <= last]
        if not hits:
            continue
        # Подпись — в конце документа: самое нижнее вхождение на самом позднем листе
        hit = max(hits, key=lambda item: (item.page, item.top, item.left))
        box = boxes[hit.page]
        middle = (hit.top + hit.bottom) / 2
        if anchor["mode"] == "cover":
            cover = (hit.left, hit.top, hit.right, hit.bottom)
            return Placement(hit.page, (hit.left + hit.right) / 2, middle, "marker", cover)
        width = float(image["widthMm"]) * MM / box.width
        height = float(image["heightMm"]) * MM / box.height
        gap = 3 * MM / box.width
        x = hit.left - gap - width / 2
        if x - width / 2 < 10 * MM / box.width:
            # Слева от фамилии места нет — над ней
            return Placement(hit.page, (hit.left + hit.right) / 2, hit.top - height / 2, "name")
        return Placement(hit.page, x, middle, "name")
    fallback = image["fallback"]
    return Placement(last, float(fallback["x"]), float(fallback["y"]), "fallback")


def image_fragment(image: dict[str, Any], placement: Placement, box: PageBox) -> str:
    """Разметка картинки на странице наложения: числа считает движок, картинка — data URI."""
    data = str(image["dataUri"])
    if not DATA_URI.match(data):
        raise RenderContentError("Картинка наложения — не PNG")
    page_width, page_height = box.width / MM, box.height / MM
    width = float(image["widthMm"])
    height = float(image["heightMm"])
    left = placement.x * page_width - width / 2
    top = placement.y * page_height - height / 2
    parts: list[str] = []
    if placement.cover:
        x0, y0, x1, y1 = placement.cover
        cover_width = (x1 - x0) * page_width + 1
        cover_height = (y1 - y0) * page_height + 1
        parts.append(
            '<div style="position:absolute;background:#fff;'
            f"left:{x0 * page_width - 0.5:.2f}mm;top:{y0 * page_height - 0.5:.2f}mm;"
            f'width:{cover_width:.2f}mm;height:{cover_height:.2f}mm"></div>'
        )
    parts.append(
        f'<img src="{data}" alt="" style="position:absolute;object-fit:contain;'
        f'left:{left:.2f}mm;top:{top:.2f}mm;width:{width:.2f}mm;height:{height:.2f}mm">'
    )
    return "".join(parts)


BLANK_OVERLAY = (
    '<!doctype html><html><head><meta charset="utf-8"><style>'
    "html,body{margin:0;padding:0;width:100%;height:100%;background:transparent}"
    "body{position:relative;overflow:hidden}</style></head><body>{body}</body></html>"
)


def with_fragments(html: str | None, fragments: list[str]) -> str:
    """Страница наложения листа: разметка api (штамп) и картинки движка поверх."""
    extra = "".join(fragments)
    if html is None:
        return BLANK_OVERLAY.replace("{body}", extra)
    marker = html.rfind("</body>")
    return html + extra if marker < 0 else html[:marker] + extra + html[marker:]


def plan_images(
    source: Path, html_pages: str, images: list[dict[str, Any]]
) -> tuple[int, dict[int, tuple[float, float]], dict[int, list[str]], list[dict[str, Any]]]:
    """Число листов, листы с наложением, фрагменты картинок по листам и места для отчёта."""
    from pypdf import PdfWriter

    writer = PdfWriter(clone_from=_open_pdf(source))
    boxes: list[PageBox] = []
    for page in writer.pages:
        page.transfer_rotation_to_content()
        box = page.cropbox
        boxes.append(
            PageBox(float(box.left), float(box.bottom), float(box.width), float(box.height))
        )
    sizes = {
        index: (boxes[index].width, boxes[index].height)
        for index in target_pages(len(boxes), html_pages)
    }
    words = page_words(source)
    fragments: dict[int, list[str]] = {}
    placements: list[dict[str, Any]] = []
    for image in images:
        placement = place_image(words, boxes, image)
        box = boxes[placement.page]
        sizes[placement.page] = (box.width, box.height)
        fragments.setdefault(placement.page, []).append(image_fragment(image, placement, box))
        placements.append(
            {
                "id": str(image["id"]),
                "page": placement.page + 1,
                "x": round(min(max(placement.x, 0.0), 1.0), 4),
                "y": round(min(max(placement.y, 0.0), 1.0), 4),
                "source": placement.source,
            }
        )
    return len(boxes), sizes, fragments, placements


def merge_page_overlays(source: Path, overlays: dict[int, bytes], target: Path) -> tuple[Path, int]:
    """Наложение по листам: у каждого листа своя страница наложения (штамп, факсимиле)."""
    from pypdf import PdfReader, PdfWriter, Transformation

    writer = PdfWriter(clone_from=_open_pdf(source))
    for index, data in sorted(overlays.items()):
        page = writer.pages[index]
        page.transfer_rotation_to_content()
        box = page.cropbox
        overlay = PdfReader(io.BytesIO(data)).pages[0]
        page.merge_transformed_page(
            overlay,
            Transformation()
            .scale(
                float(box.width) / float(overlay.mediabox.width),
                float(box.height) / float(overlay.mediabox.height),
            )
            .translate(float(box.left), float(box.bottom)),
        )
    with target.open("wb") as stream:
        writer.write(stream)
    return target, len(writer.pages)


# ─── Chromium ────────────────────────────────────────────────────────────────


async def _print_page(html: str, options: dict[str, Any]) -> bytes:
    """Страница, собранная api, — в PDF: без скриптов и без сети (все запросы отклоняются).

    Страниц печати одновременно — не больше, чем у рендера отчётов (общий семафор)."""
    async with render_slots():
        browser = await shared_browser()
        context = await browser.new_context(java_script_enabled=False)
        try:
            await context.route("**/*", lambda route: route.abort())
            page = await context.new_page()
            await page.set_content(html, wait_until="load")
            await page.emulate_media(media="print")
            data: bytes = await page.pdf(**options)
            return data
        finally:
            await context.close()


async def render_html(plan: dict[str, Any], workdir: Path) -> tuple[Path, int | None]:
    labels = plan.get("labels") or {}
    data = await _print_page(
        str(plan["html"]),
        {
            "format": "A4",
            "landscape": plan.get("orientation") == "landscape",
            "print_background": True,
            "display_header_footer": True,
            "header_template": "<span></span>",
            "footer_template": footer_template(
                str(plan.get("footer") or ""),
                str(labels.get("page") or ""),
                str(labels.get("of") or ""),
            ),
            "margin": {side: f"{value}mm" for side, value in MARGIN_MM.items()},
        },
    )
    path = workdir / "print.pdf"
    path.write_bytes(data)
    return path, count_pdf_pages(data)


def _overlay_options(width: float, height: float) -> dict[str, Any]:
    return {
        "width": f"{width / 72:.4f}in",
        "height": f"{height / 72:.4f}in",
        "print_background": True,
        "margin": {"top": "0", "right": "0", "bottom": "0", "left": "0"},
        "page_ranges": "1",
    }


async def render_overlay(
    plan: dict[str, Any], workdir: Path
) -> tuple[Path, int, list[dict[str, Any]] | None]:
    contract = document_render_contract()
    source_meta = plan["source"]
    size = await object_size(str(source_meta["bucket"]), str(source_meta["storageKey"]))
    if size > int(contract["maxSourceBytes"]):
        raise RenderContentError("Файл слишком большой для наложения")
    name = str(source_meta.get("name") or "source")
    mime = str(source_meta.get("mime") or "application/octet-stream")
    suffix = Path(name).suffix.lower() or ".bin"
    source = await download(
        str(source_meta["bucket"]), str(source_meta["storageKey"]), workdir / f"source{suffix}"
    )
    if mime != PDF_MIME and suffix != ".pdf":
        try:
            # Не-PDF (офисный файл, скан) — сначала в PDF, как PDF-представление версии
            source = await asyncio.to_thread(to_pdf, source, mime, name, workdir)
        except ToolError as error:
            raise RenderContentError(f"Файл не переводится в PDF: {error}") from error
    pages = str(plan.get("pages") or "first")
    images = list(plan.get("images") or [])
    if images:
        # Картинки (факсимиле): у каждого листа своя страница наложения
        count, sizes, fragments, placements = await asyncio.to_thread(
            plan_images, source, pages, images
        )
        html_pages = set(target_pages(count, pages))
        page_overlays: dict[int, bytes] = {}
        for index, (width, height) in sorted(sizes.items()):
            html = str(plan["html"]) if index in html_pages else None
            page_overlays[index] = await _print_page(
                with_fragments(html, fragments.get(index, [])), _overlay_options(width, height)
            )
        path, count = await asyncio.to_thread(
            merge_page_overlays, source, page_overlays, workdir / "result.pdf"
        )
        return path, count, placements
    sizes_list = await asyncio.to_thread(page_sizes, source, pages)
    overlays: dict[tuple[float, float], bytes] = {}
    for width, height in sizes_list:
        key = size_key((width, height))
        if key in overlays:
            continue
        overlays[key] = await _print_page(str(plan["html"]), _overlay_options(width, height))
    path, count = await asyncio.to_thread(
        merge_overlays, source, overlays, pages, workdir / "result.pdf"
    )
    return path, count, None


async def _template(plan: dict[str, Any], workdir: Path) -> Path:
    template = plan["template"]
    size = await object_size(str(template["bucket"]), str(template["storageKey"]))
    if size > int(document_render_contract()["maxTemplateBytes"]):
        raise RenderContentError("Шаблон слишком большой")
    return await download(
        str(template["bucket"]), str(template["storageKey"]), workdir / "template.docx"
    )


async def _produce(
    plan: dict[str, Any], workdir: Path
) -> tuple[Path, int | None, list[dict[str, Any]] | None]:
    kind = plan["kind"]
    if kind == "html":
        path, pages = await render_html(plan, workdir)
        return path, pages, None
    if kind == "overlay":
        return await render_overlay(plan, workdir)
    if kind == "docx":
        source = await _template(plan, workdir)
        context = plan.get("context") or {}
        filled = await asyncio.to_thread(fill_template, source, context, workdir / "filled.docx")
        return filled, None, None
    raise RenderContentError(f"Неизвестный вид рендера: {kind}")


@handler("render", "document.render")
async def document_render(data: dict[str, Any]) -> dict[str, Any]:
    render_id = str(data["renderId"])
    started = time.monotonic()
    start = await document_render_start(render_id)
    if start.get("status") != "render":
        log.info("document.render_skipped", render_id=render_id, reason=start.get("reason"))
        return {"skipped": start.get("reason")}
    plan: dict[str, Any] = start["plan"]
    target: dict[str, Any] | None = start.get("target")

    with tempfile.TemporaryDirectory(prefix="kchs-render-") as tmp:
        workdir = Path(tmp)
        try:
            if plan["kind"] == "inspect":
                source = await _template(plan, workdir)
                placeholders = await asyncio.to_thread(inspect_template, source)
                await document_render_done(
                    render_id, {"status": "ready", "placeholders": placeholders}
                )
                return {"placeholders": len(placeholders)}
            output, pages, placements = await asyncio.wait_for(
                _produce(plan, workdir), timeout=settings().ENGINE_RENDER_TIMEOUT_S
            )
        except RenderContentError as error:
            await document_render_done(render_id, {"status": "failed", "error": str(error)[:4000]})
            log.warning("document.render_failed", render_id=render_id, error=str(error))
            return {"failed": str(error)}
        except TimeoutError as error:
            raise RuntimeError("рендер не уложился в отведённое время") from error
        if target is None:
            raise RuntimeError("api не выдало ключ результата")
        await upload(
            str(target["bucket"]), str(target["storageKey"]), output, str(target["contentType"])
        )
        size = output.stat().st_size

    done: dict[str, Any] = {"status": "ready", "size": size, "pages": pages}
    if placements is not None:
        done["placements"] = placements
    await document_render_done(render_id, done)
    duration = int((time.monotonic() - started) * 1000)
    log.info(
        "document.rendered",
        render_id=render_id,
        kind=plan["kind"],
        pages=pages,
        duration_ms=duration,
    )
    return {"kind": plan["kind"], "pages": pages, "size": size, "durationMs": duration}
