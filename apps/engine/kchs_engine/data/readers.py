"""Чтение файлов импорта: формат, кодировка, разделитель, листы и строки.

Источник отдаёт записи с начала файла парами «номер строки, ячейки»; номер —
такой, как его видит человек: номер записи CSV (строка заголовка — 1), номер
строки листа Excel, номер объекта JSON или строки NDJSON.

Анализ видит «голову» файла (`open_head`): начало текстового файла до
`SAMPLE_BYTES` или книгу Excel. Нормализация строит ту же голову из всего файла
— строки над заголовком, десятичный знак, порядок частей даты и столбцы JSON
определяются по тем же данным, что и при анализе, — и затем читает весь файл
потоком (`open_full`). Поэтому номера строк и индексы столбцов у анализа и
нормализации совпадают.
"""

import codecs
import csv
import io
import json
import sys
from collections import Counter
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any, BinaryIO
from zipfile import BadZipFile

# Сколько байт текстового файла видит анализ (голова файла)
SAMPLE_BYTES = 8 * 1024 * 1024
# JSON целиком в памяти: больше — только NDJSON или CSV
MAX_JSON_BYTES = 256 * 1024 * 1024
# Длинные ячейки CSV (многострочный текст)
csv.field_size_limit(min(sys.maxsize, 64 * 1024 * 1024))

DELIMITERS = (";", ",", "\t", "|")
JSON_FORMATS = frozenset({"json", "ndjson", "geojson"})
EXCEL_FORMATS = frozenset({"xlsx", "xls"})


class ImportFileError(Exception):
    """Файл нельзя прочитать: повтор задания не поможет.

    `code` — для API и тестов, `message` — причина по-русски со строчной буквы
    (её дополняют: «Не удалось прочитать файл: …»).
    """

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


# ─── Кодировка ───────────────────────────────────────────────────────────────


def _fallback_cp1251(error: UnicodeError) -> tuple[str, int]:
    """Байт не из UTF-8 посреди UTF-8-файла — чаще всего кириллица Windows-1251.

    Такие файлы получаются, когда в выгрузку UTF-8 дописывают строки из старой
    программы; начало файла при этом может быть чистым ASCII.
    """
    if not isinstance(error, UnicodeDecodeError):
        raise error
    chunk = bytes(error.object[error.start : error.end])
    return chunk.decode("cp1251", errors="replace"), error.end


codecs.register_error("kchs-cp1251", _fallback_cp1251)


def decode_errors(encoding: str) -> str:
    """Обработчик ошибок декодирования: для UTF-8 — запасная Windows-1251."""
    name = codecs.lookup(encoding).name
    return "kchs-cp1251" if name in ("utf-8", "utf-8-sig") else "replace"


# ─── Формат и кодировка ──────────────────────────────────────────────────────

_EXTENSIONS = {
    ".csv": "csv",
    ".txt": "csv",
    ".tsv": "tsv",
    ".tab": "tsv",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".xls": "xls",
    ".json": "json",
    ".ndjson": "ndjson",
    ".jsonl": "ndjson",
    ".geojson": "geojson",
}


def detect_format(head: bytes, file_name: str, hint: str | None = None) -> str:
    """Формат по содержимому, затем по расширению: .csv с книгой Excel внутри — всё равно Excel."""
    if hint:
        return hint
    if head.startswith(b"PK\x03\x04"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"
    suffix = Path(file_name).suffix.lower()
    by_name = _EXTENSIONS.get(suffix)
    text = _decode_head(head)
    stripped = text.lstrip()
    if stripped.startswith("{") or stripped.startswith("["):
        if '"FeatureCollection"' in stripped[:4096] or by_name == "geojson":
            return "geojson"
        if stripped.startswith("{") and "\n" in stripped:
            lines = [line.strip() for line in stripped.splitlines()[:5] if line.strip()]
            if len(lines) > 1 and all(
                line.startswith("{") and line.endswith("}") for line in lines
            ):
                return "ndjson"
        if by_name in ("json", "ndjson", "geojson"):
            return by_name
        return "json"
    if by_name in ("csv", "tsv"):
        return by_name
    return by_name or "csv"


def _decode_head(head: bytes) -> str:
    sample = head[:65536]
    return sample.decode(detect_encoding(sample), errors="replace").lstrip("\ufeff")


# Частые строчные буквы русского текста — по ним отличаем верную кириллическую кодировку
_FREQUENT = set("оеаинтсрвлкмдпуяыь")
_CYRILLIC_CANDIDATES = ("cp1251", "koi8_r", "cp866")
ENCODING_NAMES = {
    "utf-8": "UTF-8",
    "utf-8-sig": "UTF-8 (BOM)",
    "utf-16": "UTF-16",
    "utf-16-le": "UTF-16LE",
    "utf-16-be": "UTF-16BE",
    "cp1251": "Windows-1251",
    "koi8_r": "KOI8-R",
    "cp866": "CP866 (DOS)",
    "cp1252": "Windows-1252",
}


def _cyrillic_score(text: str) -> float:
    letters = [char for char in text if char.isalpha() and "Ѐ" <= char <= "ӿ"]
    if len(letters) < 20:
        return 0.0
    return sum(1 for char in letters if char in _FREQUENT) / len(letters)


def detect_encoding(sample: bytes) -> str:
    """UTF-8 (с BOM и без), UTF-16, иначе однобайтная кириллица по частоте букв."""
    if sample.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    if len(sample) >= 16:
        even, odd = sample[0::2], sample[1::2]
        even_nul, odd_nul = even.count(0) / len(even), odd.count(0) / len(odd)
        if odd_nul > 0.3 and even_nul < 0.05:
            return "utf-16-le"
        if even_nul > 0.3 and odd_nul < 0.05:
            return "utf-16-be"
    try:
        sample.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError as error:
        # Выборка оборвалась посреди многобайтного символа — это всё ещё UTF-8
        if error.start >= len(sample) - 3:
            try:
                sample[: error.start].decode("utf-8")
                return "utf-8"
            except UnicodeDecodeError:
                pass
    scores = {
        name: _cyrillic_score(sample.decode(name, errors="replace"))
        for name in _CYRILLIC_CANDIDATES
    }
    best = max(scores, key=lambda name: scores[name])
    if scores[best] >= 0.3:
        return best
    return "cp1252"


def detect_delimiter(text: str, fmt: str) -> str:
    """Разделитель, дающий одинаковое число полей в большинстве строк.

    Строки-заголовки отчёта над таблицей не мешают: считается самая частая
    ширина строки, а не ширина первой.
    """
    if fmt == "tsv":
        return "\t"
    lines = [line for line in text.splitlines()[:300] if line.strip()]
    if not lines:
        return ","
    sample = "\n".join(lines)
    best, best_score = ",", (0, 0)
    for candidate in DELIMITERS:
        try:
            counts = [len(row) for row in csv.reader(io.StringIO(sample), delimiter=candidate)]
        except csv.Error:
            continue
        common = Counter(count for count in counts if count > 1).most_common(1)
        if not common:
            continue
        width, frequency = common[0]
        score = (frequency, width)
        if score > best_score:
            best, best_score = candidate, score
    return best


# ─── Источники строк ─────────────────────────────────────────────────────────

Row = tuple[int, list[Any]]


class Source:
    """Записи файла с начала: (номер строки как в файле, ячейки)."""

    format: str = "csv"

    def rows(self) -> Iterator[Row]:
        raise NotImplementedError

    def progress(self) -> float:
        """Доля прочитанного, 0…1 — для отчёта о ходе задания."""
        return 0.0

    def close(self) -> None:
        return None


class _CountingReader(io.RawIOBase):
    """Поток байтов со счётчиком прочитанного — доля файла для прогресса."""

    def __init__(self, raw: BinaryIO) -> None:
        self.raw = raw
        self.count = 0

    def readable(self) -> bool:
        return True

    def readinto(self, buffer: Any) -> int:
        data = self.raw.read(len(buffer))
        size = len(data)
        buffer[:size] = data
        self.count += size
        return size


class DelimitedSource(Source):
    """CSV/TSV: весь файл потоком (нормализация)."""

    def __init__(self, path: Path, encoding: str, delimiter: str, fmt: str = "csv") -> None:
        self.format = fmt
        self.path = path
        self.encoding = encoding
        self.delimiter = delimiter
        self.size = max(path.stat().st_size, 1)
        self._counter: _CountingReader | None = None

    def rows(self) -> Iterator[Row]:
        with self.path.open("rb") as raw:
            self._counter = _CountingReader(raw)
            buffered = io.BufferedReader(self._counter, buffer_size=1 << 20)
            stream = io.TextIOWrapper(
                buffered,
                encoding=self.encoding,
                errors=decode_errors(self.encoding),
                newline="",
            )
            lines = (line.replace("\x00", "") if "\x00" in line else line for line in stream)
            reader = csv.reader(lines, delimiter=self.delimiter)
            number = 0
            try:
                for record in reader:
                    number += 1
                    if number == 1 and record and record[0].startswith("\ufeff"):
                        record[0] = record[0][1:]
                    yield number, record
            except csv.Error as error:
                raise ImportFileError("unreadable", f"запись {number + 1}: {error}") from error

    def progress(self) -> float:
        if self._counter is None:
            return 0.0
        return min(self._counter.count / self.size, 1.0)


def _newline_bytes(encoding: str, data: bytes) -> bytes:
    """Перевод строки в байтах кодировки (UTF-16 — два байта, порядок по BOM)."""
    name = codecs.lookup(encoding).name
    if name == "utf-16-be" or (name == "utf-16" and data.startswith(codecs.BOM_UTF16_BE)):
        return b"\x00\n"
    if name.startswith("utf-16"):
        return b"\n\x00"
    return b"\n"


def sample_text(path: Path, encoding: str, complete: bool) -> tuple[str, bool, int]:
    """Голова текстового файла: текст, признак «весь файл» и сколько байт она покрывает.

    Если файл длиннее `SAMPLE_BYTES` (или у нас только его начало), неполная
    последняя строка отбрасывается.
    """
    with path.open("rb") as stream:
        data = stream.read(SAMPLE_BYTES + 1)
    whole = complete and len(data) <= SAMPLE_BYTES
    data = data[:SAMPLE_BYTES]
    if not whole:
        newline = _newline_bytes(encoding, data)
        cut = data.rfind(newline)
        if cut > 0:
            data = data[: cut + len(newline)]
    text = data.decode(encoding, errors=decode_errors(encoding)).lstrip("\ufeff")
    return text.replace("\x00", ""), whole, len(data)


class TextSampleSource(Source):
    """Записи CSV из головы файла (анализ и параметры нормализации)."""

    def __init__(self, text: str, delimiter: str, fmt: str = "csv") -> None:
        self.format = fmt
        self.text = text
        self.delimiter = delimiter

    def rows(self) -> Iterator[Row]:
        reader = csv.reader(io.StringIO(self.text, newline=""), delimiter=self.delimiter)
        number = 0
        try:
            for record in reader:
                number += 1
                yield number, record
        except csv.Error:
            # Оборванная кавычка в конце выборки — анализу хватит прочитанного
            return

    def count_rows(self, start: int) -> int:
        """Непустые записи головы после строки `start` — для оценки числа строк файла."""
        count = 0
        for number, record in self.rows():
            if number > start and any(cell.strip() for cell in record):
                count += 1
        return count


def _max_row(worksheet: Any) -> int:
    """Число строк листа по его размерам (у книг без размеров — 0)."""
    try:
        return int(worksheet.max_row or 0)
    except (TypeError, ValueError):
        return 0


class XlsxSource(Source):
    def __init__(self, path: Path, sheet: str | None) -> None:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        self.format = "xlsx"
        # Файл, а не путь: openpyxl проверяет расширение имени, а временный файл
        # задания и анализа называется без него
        self._stream = path.open("rb")
        try:
            self.book = load_workbook(self._stream, read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError) as error:
            self._stream.close()
            raise ImportFileError(
                "unreadable", "книга Excel повреждена или защищена паролем"
            ) from error
        self.sheets = [(str(ws.title), _max_row(ws)) for ws in self.book.worksheets]
        for worksheet in self.book.worksheets:
            # Размеры листа из файла бывают неверными (A1 у выгрузок сторонних
            # программ) — строки читаются по фактическому содержимому
            if hasattr(worksheet, "reset_dimensions"):
                worksheet.reset_dimensions()
        self.sheet = sheet or self._first_with_data()
        if self.sheet not in self.book.sheetnames:
            raise ImportFileError("sheet_not_found", f"в книге нет листа «{self.sheet}»")
        self.worksheet = self.book[self.sheet]
        self.total = max(dict(self.sheets).get(self.sheet, 0), 1)
        self._done = 0

    def _first_with_data(self) -> str:
        """Первый лист, где в первых строках есть строка хотя бы из двух значений."""
        for worksheet in self.book.worksheets:
            for row in worksheet.iter_rows(max_row=50, values_only=True):
                if sum(1 for value in row if value not in (None, "")) >= 2:
                    return str(worksheet.title)
        return str(self.book.worksheets[0].title) if self.book.worksheets else ""

    def rows(self) -> Iterator[Row]:
        for number, row in enumerate(self.worksheet.iter_rows(values_only=True), start=1):
            self._done = number
            yield number, list(row)

    def progress(self) -> float:
        return min(self._done / self.total, 1.0)

    def close(self) -> None:
        self.book.close()
        self._stream.close()


class XlsSource(Source):
    """Старый формат Excel 97–2003 (xlrd)."""

    def __init__(self, path: Path, sheet: str | None) -> None:
        import xlrd

        self.format = "xls"
        self._xlrd = xlrd
        try:
            self.book = xlrd.open_workbook(str(path), on_demand=True)
        except Exception as error:  # xlrd бросает свои исключения без общего предка
            raise ImportFileError("unreadable", "книга Excel 97–2003 повреждена") from error
        self.sheets: list[tuple[str, int]] = []
        for index in range(self.book.nsheets):
            worksheet = self.book.sheet_by_index(index)
            self.sheets.append((str(worksheet.name), int(worksheet.nrows)))
        names = [name for name, _rows in self.sheets]
        chosen = sheet or self._first_with_data(names)
        if chosen not in names:
            raise ImportFileError("sheet_not_found", f"в книге нет листа «{chosen}»")
        self.sheet = chosen
        self.worksheet = self.book.sheet_by_name(chosen)
        self.total = max(int(self.worksheet.nrows), 1)
        self._done = 0

    def _first_with_data(self, names: list[str]) -> str:
        for name in names:
            worksheet = self.book.sheet_by_name(name)
            for index in range(min(worksheet.nrows, 50)):
                values = [self._value(cell) for cell in worksheet.row(index)]
                if sum(1 for value in values if value not in (None, "")) >= 2:
                    return name
        return names[0] if names else ""

    def _value(self, cell: Any) -> Any:
        xlrd = self._xlrd
        if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
            return None
        if cell.ctype == xlrd.XL_CELL_DATE:
            try:
                return xlrd.xldate.xldate_as_datetime(cell.value, self.book.datemode)
            except (ValueError, OverflowError, xlrd.xldate.XLDateError):
                return cell.value
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        if cell.ctype == xlrd.XL_CELL_NUMBER:
            value = float(cell.value)
            return int(value) if value.is_integer() and abs(value) < 1e15 else value
        return cell.value

    def rows(self) -> Iterator[Row]:
        for index in range(self.worksheet.nrows):
            self._done = index + 1
            yield index + 1, [self._value(cell) for cell in self.worksheet.row(index)]

    def progress(self) -> float:
        return min(self._done / self.total, 1.0)

    def close(self) -> None:
        self.book.release_resources()


# ─── JSON, NDJSON, GeoJSON ───────────────────────────────────────────────────

_DECODER = json.JSONDecoder()
_WRAPPER_KEYS = ("data", "items", "rows", "records", "results", "features")

Located = tuple[Any, int]


class BrokenLine(str):
    """Строка NDJSON, которая не читается как JSON."""


def _iter_array(text: str, start: int, complete: bool) -> Iterator[Located]:
    """Элементы JSON-массива с позиции `[` и позиция конца каждого.

    Голова файла обрывается посреди объекта — это просто конец выборки; в
    целом файле та же ошибка — файл не читается.
    """
    position = start + 1
    length = len(text)
    while position < length:
        while position < length and text[position] in " \t\r\n,":
            position += 1
        if position >= length or text[position] == "]":
            return
        try:
            value, position = _DECODER.raw_decode(text, position)
        except json.JSONDecodeError as error:
            if complete:
                raise ImportFileError(
                    "unreadable",
                    f"JSON не читается: строка {error.lineno}, позиция {error.colno}",
                ) from error
            return
        yield value, position


def _records_from_text(text: str, fmt: str, complete: bool) -> tuple[Iterator[Located], list[str]]:
    """Объекты JSON/GeoJSON из текста (целого или головы файла) и предупреждения."""
    warnings: list[str] = []
    stripped = text.lstrip()
    offset = len(text) - len(stripped)
    if fmt == "geojson":
        marker = text.find('"features"')
        bracket = text.find("[", marker) if marker >= 0 else -1
        if bracket < 0:
            raise ImportFileError("unreadable", "в GeoJSON нет списка features")
        crs = text.find('"crs"', 0, marker)
        if crs >= 0:
            declared = text[crs : crs + 200]
            if "4326" not in declared and "CRS84" not in declared:
                warnings.append(
                    "В GeoJSON указана система координат, отличная от WGS 84: "
                    "геометрия загружается только в EPSG:4326, пересчёт не выполняется"
                )
        return _iter_array(text, bracket, complete), warnings
    if stripped.startswith("["):
        return _iter_array(text, offset, complete), warnings
    if stripped.startswith("{"):
        # Обёртка {"data": [...]}: берём первый список объектов
        for key in _WRAPPER_KEYS:
            marker = text.find(f'"{key}"')
            if marker >= 0:
                bracket = text.find("[", marker)
                between = text[marker + len(key) + 2 : bracket] if bracket >= 0 else ""
                if bracket >= 0 and between.strip() in (":", ""):
                    return _iter_array(text, bracket, complete), warnings
        try:
            value, end = _DECODER.raw_decode(text, offset)
        except json.JSONDecodeError as error:
            raise ImportFileError("unreadable", "файл не читается как JSON") from error
        if isinstance(value, dict):
            return iter([(value, end)]), warnings
    raise ImportFileError("unreadable", "в JSON нет массива объектов")


class JsonSource(Source):
    """JSON-массив объектов, NDJSON, GeoJSON FeatureCollection.

    Голова (`text`) — анализ; весь файл (`path`) — нормализация. Столбцы — ключи
    объектов (у GeoJSON — свойства) в порядке первого появления в первых
    объектах головы (`discover`); у GeoJSON геометрия объекта — последняя,
    служебная ячейка строки (`geometry_index`).
    """

    def __init__(
        self,
        fmt: str,
        *,
        text: str | None = None,
        path: Path | None = None,
        encoding: str = "utf-8",
        whole: bool = True,
        columns: list[str] | None = None,
    ) -> None:
        self.format = fmt
        self.path = path
        self.encoding = encoding
        self.whole = whole
        self.warnings: list[str] = []
        self._text = text
        self._progress = 0.0
        self.columns: list[str] = list(columns or [])
        self.geometry_index: int | None = len(self.columns) if fmt == "geojson" else None

    def _load_text(self) -> str:
        if self._text is None:
            if self.path is None:  # pragma: no cover — источник без текста и без файла
                raise ImportFileError("unreadable", "нет данных JSON")
            if self.path.stat().st_size > MAX_JSON_BYTES:
                raise ImportFileError(
                    "too_large",
                    "JSON больше 256 МБ читается только целиком — "
                    "сохраните данные в NDJSON или CSV",
                )
            raw = self.path.read_text(encoding=self.encoding, errors=decode_errors(self.encoding))
            self._text = raw.lstrip("\ufeff")
        return self._text

    def _objects(self) -> Iterator[tuple[int, Any]]:
        if self.format == "ndjson":
            yield from self._lines()
            return
        text = self._load_text()
        records, warnings = _records_from_text(text, self.format, self.whole)
        for warning in warnings:
            if warning not in self.warnings:
                self.warnings.append(warning)
        total = max(len(text), 1)
        for number, (record, end) in enumerate(records, start=1):
            self._progress = end / total
            yield number, record

    def _lines(self) -> Iterator[tuple[int, Any]]:
        counter: _CountingReader | None = None
        raw: BinaryIO | None = None
        stream: io.TextIOBase
        if self._text is not None:
            # Те же правила перевода строк, что у файла: \n, \r\n и \r
            stream = io.StringIO(self._text, newline=None)
            size = 1
        else:
            if self.path is None:  # pragma: no cover
                raise ImportFileError("unreadable", "нет данных NDJSON")
            raw = self.path.open("rb")
            counter = _CountingReader(raw)
            size = max(self.path.stat().st_size, 1)
            stream = io.TextIOWrapper(
                io.BufferedReader(counter, buffer_size=1 << 20),
                encoding=self.encoding,
                errors=decode_errors(self.encoding),
                newline=None,
            )
        try:
            for number, line in enumerate(stream, start=1):
                if counter is not None:
                    self._progress = min(counter.count / size, 1.0)
                text = line.strip().lstrip("\ufeff")
                if not text:
                    continue
                try:
                    yield number, json.loads(text)
                except ValueError:
                    yield number, BrokenLine(text)
        finally:
            if raw is not None:
                raw.close()

    def _properties(self, record: Any) -> Any:
        if self.format == "geojson" and isinstance(record, dict):
            properties = record.get("properties")
            return properties if isinstance(properties, dict) else {}
        return record

    def discover(self, limit: int) -> None:
        """Столбцы — ключи первых `limit` объектов в порядке появления."""
        columns: dict[str, None] = {}
        for index, (_number, record) in enumerate(self._objects()):
            if index >= limit:
                break
            properties = self._properties(record)
            if isinstance(properties, dict):
                for key in properties:
                    columns.setdefault(str(key), None)
        self.columns = list(columns)
        self.geometry_index = len(self.columns) if self.format == "geojson" else None
        self._progress = 0.0

    def rows(self) -> Iterator[Row]:
        index = {name: position for position, name in enumerate(self.columns)}
        width = len(self.columns)
        for number, record in self._objects():
            properties = self._properties(record)
            cells: list[Any] = [None] * width
            if isinstance(properties, dict):
                for key, value in properties.items():
                    position = index.get(str(key))
                    if position is not None:
                        cells[position] = value
            elif width:
                # Не объект (число, строка, массив) — значение первого столбца
                cells[0] = properties
            else:
                cells = [properties]
            if self.geometry_index is not None:
                geometry = record.get("geometry") if isinstance(record, dict) else None
                cells.append(geometry)
            yield number, cells

    def count_rows(self) -> tuple[int, float]:
        """Объекты головы и доля головы, которую они занимают (для оценки числа строк)."""
        count = 0
        for _number, _record in self._objects():
            count += 1
        return count, self._progress if self.format != "ndjson" else 1.0

    def progress(self) -> float:
        return self._progress


# ─── Открытие ────────────────────────────────────────────────────────────────


@dataclass
class Opened:
    """Голова файла и распознанные параметры чтения."""

    format: str
    encoding: str | None
    delimiter: str | None
    source: Source
    sheets: list[tuple[str, int]]
    sheet: str | None
    warnings: list[str]
    # Голова — это весь файл; сколько байт файла она покрывает
    whole: bool
    head_bytes: int


def _read_head(path: Path, size: int = 1 << 20) -> bytes:
    with path.open("rb") as stream:
        return stream.read(size)


def open_head(
    path: Path, file_name: str, options: dict[str, Any], *, complete: bool, discover: int
) -> Opened:
    """Голова файла по параметрам (пустые — автоопределение).

    `complete` — в `path` весь файл, а не только его начало. Книгу Excel можно
    открыть только целиком.
    """
    head = _read_head(path, 65536)
    if not head:
        raise ImportFileError("empty", "файл пуст")
    fmt = detect_format(head, file_name, options.get("format"))
    if fmt in EXCEL_FORMATS:
        if not complete:  # pragma: no cover — анализ скачивает книгу целиком
            raise ImportFileError("unreadable", "книгу Excel нельзя прочитать по началу файла")
        sheet = options.get("sheet")
        book: XlsxSource | XlsSource = (
            XlsxSource(path, sheet) if fmt == "xlsx" else XlsSource(path, sheet)
        )
        size = path.stat().st_size
        return Opened(fmt, None, None, book, book.sheets, book.sheet, [], True, size)

    encoding = options.get("encoding") or detect_encoding(_read_head(path))
    try:
        codecs.lookup(encoding)
    except LookupError as error:
        raise ImportFileError("unsupported", f"кодировка {encoding} не поддерживается") from error
    text, whole, head_bytes = sample_text(path, encoding, complete)

    if fmt in JSON_FORMATS:
        json_source = JsonSource(fmt, text=text, encoding=encoding, whole=whole)
        json_source.discover(discover)
        return Opened(
            fmt, encoding, None, json_source, [], None, json_source.warnings, whole, head_bytes
        )

    delimiter = options.get("delimiter") or detect_delimiter(text, fmt)
    return Opened(
        fmt,
        encoding,
        delimiter,
        TextSampleSource(text, delimiter, fmt),
        [],
        None,
        [],
        whole,
        head_bytes,
    )


def open_full(path: Path, head: Opened) -> Source:
    """Весь файл потоком с параметрами головы (нормализация)."""
    if head.format in EXCEL_FORMATS:
        # Строки книги читаются заново с первой — те же, что видела голова
        return head.source
    encoding = head.encoding or "utf-8"
    if head.format in JSON_FORMATS:
        columns = head.source.columns if isinstance(head.source, JsonSource) else []
        return JsonSource(head.format, path=path, encoding=encoding, columns=columns)
    return DelimitedSource(path, encoding, head.delimiter or ",", head.format)
