"""Импорт пользователей из Excel (P0-E04 S04, ADR-0041): разбор XLSX и шаблон.

Движок только читает файл и приводит ячейки к тексту. Проверка строк
(логин, почта, роли, подразделения, дубликаты) и создание пользователей —
в API: там оргструктура, роли и права администратора. Столбцы и коды
замечаний — контракт `users_import.json` из `packages/contracts`.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from kchs_engine.contracts import users_import_contract

# Заголовок ищется в первых строках листа: над ним бывает название таблицы
HEADER_SEARCH_ROWS = 20
# Столько пустых строк подряд — конец данных (лист с оформлением «до конца»)
EMPTY_ROWS_STOP = 1000
# Значение ячейки длиннее этого приходит обрезанным: пределы полей меньше,
# и API всё равно отметит строку как слишком длинную
MAX_CELL_CHARS = 1000

LOCALES = ("ru", "tg", "en")


def normalize_header(value: object) -> str:
    """Заголовок для сравнения: регистр, «ё», пробелы и пометка `*` не важны."""
    text = str(value).replace("ё", "е").replace("Ё", "Е").lower().replace("*", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip(" .:;")


def _header_index() -> dict[str, str]:
    """Нормализованный заголовок → ключ поля (заголовки на всех языках и синонимы)."""
    index: dict[str, str] = {}
    for spec in users_import_contract()["fields"]:
        names = [*spec["headers"].values(), *spec.get("aliases", []), spec["key"]]
        for name in names:
            index.setdefault(normalize_header(name), spec["key"])
    return index


def cell_text(value: object) -> str | None:
    """Ячейка → текст так, как его видит человек в Excel.

    Число без дробной части — без «.0» (телефон 992935001122 хранится как
    float), даты — ISO, пустые строки и пробелы — `None`.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.date().isoformat()
        return value.isoformat(timespec="minutes")
    if isinstance(value, date | time):
        return value.isoformat()
    text = str(value).replace(" ", " ").strip()
    if not text:
        return None
    return text[:MAX_CELL_CHARS]


@dataclass
class ParseResult:
    rows: list[dict[str, Any]] = field(default_factory=list)
    columns: dict[str, str] = field(default_factory=dict)
    warnings: list[dict[str, Any]] = field(default_factory=list)
    file_error: dict[str, Any] | None = None
    total_rows: int = 0

    def payload(self) -> dict[str, Any]:
        return {
            "rows": self.rows,
            "columns": self.columns,
            "warnings": self.warnings,
            "fileError": self.file_error,
            "totalRows": self.total_rows,
        }


def _issue(code: str, **params: str | int) -> dict[str, Any]:
    return {"code": code, "field": None, "params": params}


def _failure(code: str, **params: str | int) -> ParseResult:
    return ParseResult(file_error=_issue(code, **params))


@dataclass
class _Header:
    row: int
    mapping: dict[int, str]  # номер столбца → поле
    columns: dict[str, str]  # поле → заголовок в файле
    unknown: list[str]
    duplicate: str | None


def _find_header(rows: list[tuple[Any, ...]], index: dict[str, str]) -> _Header | None:
    for offset, values in enumerate(rows):
        mapping: dict[int, str] = {}
        columns: dict[str, str] = {}
        unknown: list[str] = []
        duplicate: str | None = None
        for column, raw in enumerate(values):
            title = cell_text(raw)
            if title is None:
                continue
            key = index.get(normalize_header(title))
            if key is None:
                unknown.append(title)
            elif key in columns:
                duplicate = duplicate or title
            else:
                mapping[column] = key
                columns[key] = title
        if mapping:
            return _Header(offset + 1, mapping, columns, unknown, duplicate)
    return None


def parse_workbook(path: Path) -> ParseResult:
    """Разбор файла импорта. Ошибка формата — замечание к файлу, а не исключение:
    повтор задания её не исправит."""
    contract = users_import_contract()
    max_rows = int(contract["maxRows"])
    index = _header_index()
    required = [spec for spec in contract["fields"] if spec["required"]]

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError) as error:
        return _failure("unreadable", reason=str(error)[:200])

    try:
        sheets = [sheet for sheet in workbook.worksheets if hasattr(sheet, "iter_rows")]
        if not sheets:
            return _failure("no_sheet")

        # Лист с данными — первый, где узнаётся заголовок (в шаблоне есть справочные листы)
        chosen = None
        header: _Header | None = None
        for sheet in sheets:
            top = list(sheet.iter_rows(max_row=HEADER_SEARCH_ROWS, values_only=True))
            found = _find_header(top, index)
            if found and all(spec["key"] in found.columns for spec in required):
                chosen, header = sheet, found
                break
            if found and header is None:
                chosen, header = sheet, found
        if chosen is None or header is None:
            return _failure("no_header")

        if header.duplicate:
            return _failure("duplicate_column", column=header.duplicate)
        missing = [spec["headers"]["ru"] for spec in required if spec["key"] not in header.columns]
        if missing:
            return _failure("missing_columns", columns=", ".join(missing))

        result = ParseResult(columns=header.columns)
        result.warnings = [_issue("unknown_column", column=title) for title in header.unknown]

        empty_streak = 0
        for number, values in enumerate(
            chosen.iter_rows(min_row=header.row + 1, values_only=True), start=header.row + 1
        ):
            record = {
                key: cell_text(values[column]) if column < len(values) else None
                for column, key in header.mapping.items()
            }
            if all(value is None for value in record.values()):
                empty_streak += 1
                if empty_streak >= EMPTY_ROWS_STOP:
                    break
                continue
            empty_streak = 0
            result.total_rows += 1
            if result.total_rows > max_rows:
                return _failure("too_many_rows", max=max_rows)
            result.rows.append({"row": number, "values": record})
        return result
    finally:
        workbook.close()


# ── Шаблон ───────────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="E9EFFF")
REQUIRED_FONT = Font(bold=True)
HINTS: dict[str, str] = {
    "login": "Латиница, цифры, точка, дефис и подчёркивание; 3–64 символа. Обязательно.",
    "lastName": "Обязательно.",
    "firstName": "Обязательно.",
    "middleName": "",
    "email": "Уникальна в системе. Нужна для восстановления пароля по ссылке.",
    "phone": "В любом формате, до 32 символов.",
    "unit": "Код или точное название подразделения — см. лист «Подразделения».",
    "position": "Название должности — см. лист «Должности». Только вместе с подразделением.",
    "roles": "Ключи или названия ролей через запятую — см. лист «Роли». Пусто — «Сотрудник».",
    "locale": "ru, tg или en. Пусто — ru.",
    "timezone": "Например, Asia/Dushanbe. Пусто — Asia/Dushanbe.",
}


def _sized(sheet: Any, widths: list[int]) -> None:
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def build_template(
    target: Path,
    roles: list[dict[str, str]],
    units: list[dict[str, str]],
    positions: list[dict[str, str]],
) -> Path:
    """Шаблон импорта: лист данных с русскими заголовками и справочные листы
    с текущими ролями, подразделениями и должностями организации."""
    contract = users_import_contract()
    fields = contract["fields"]
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Пользователи"
    for column, spec in enumerate(fields, start=1):
        title = spec["headers"]["ru"] + (" *" if spec["required"] else "")
        cell = sheet.cell(row=1, column=column, value=title)
        cell.fill = HEADER_FILL
        cell.font = REQUIRED_FONT if spec["required"] else Font()
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    _sized(sheet, [max(14, len(spec["headers"]["ru"]) + 6) for spec in fields])
    # Логин и телефон — текст: Excel не превратит «0012» в число и не отрежет нули
    for key in ("login", "phone"):
        column = next(i for i, spec in enumerate(fields, start=1) if spec["key"] == key)
        sheet.column_dimensions[get_column_letter(column)].number_format = "@"
    locale_column = next(i for i, spec in enumerate(fields, start=1) if spec["key"] == "locale")
    choice = DataValidation(type="list", formula1=f'"{",".join(LOCALES)}"', allow_blank=True)
    letter = get_column_letter(locale_column)
    choice.add(f"{letter}2:{letter}{int(contract['maxRows']) + 1}")
    sheet.add_data_validation(choice)

    guide = workbook.create_sheet("Справка")
    guide.append(["Столбец", "Как заполнять"])
    for spec in fields:
        title = spec["headers"]["ru"] + (" *" if spec["required"] else "")
        guide.append([title, HINTS.get(spec["key"], "")])
    guide.append([])
    guide.append(
        [
            "Ограничения",
            f"Не больше {contract['maxRows']} строк и "
            f"{int(contract['maxBytes']) // (1024 * 1024)} МБ. "
            "Строки с уже существующим логином пропускаются — повторный импорт не создаёт дублей.",
        ]
    )
    guide.append(
        [
            "Доступ",
            "Созданные получают временный пароль (файл выдаётся администратору один раз) "
            "и при первом входе задают свой.",
        ]
    )
    for cell in guide[1]:
        cell.font = REQUIRED_FONT
    _sized(guide, [26, 90])

    reference = [
        ("Роли", ["Ключ", "Название"], [[r["key"], r["name"]] for r in roles]),
        (
            "Подразделения",
            ["Код", "Название", "Действует"],
            [[u["code"], u["name"], "да" if u.get("active", True) else "нет"] for u in units],
        ),
        (
            "Должности",
            ["Название", "Подразделение (код)"],
            [[p["name"], p.get("unitCode") or ""] for p in positions],
        ),
    ]
    for title, header, rows in reference:
        extra = workbook.create_sheet(title)
        extra.append(header)
        for cell in extra[1]:
            cell.font = REQUIRED_FONT
        for row in rows:
            extra.append(row)
        _sized(extra, [24, 60, 12][: len(header)])

    workbook.save(target)
    return target
