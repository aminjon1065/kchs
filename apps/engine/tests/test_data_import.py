"""Импорт датасетов (P1-E02 S01, S03; ADR-0046): анализ и нормализация «грязных» файлов.

Каждый файл создаётся в тесте: кодировки, разделители, строки над заголовком,
форматы чисел и дат, книги Excel, JSON, GeoJSON, геометрия и ошибки в данных.
"""

import json
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest
from import_helpers import analyze, column, mapping_from, normalize
from openpyxl import Workbook

from kchs_engine.data.analyze import analyze_file
from kchs_engine.data.readers import SAMPLE_BYTES, ImportFileError

EXCEL_EPOCH = date(1899, 12, 30)


def write(path: Path, text: str, encoding: str = "utf-8") -> Path:
    path.write_bytes(text.encode(encoding))
    return path


def keys(analysis: dict) -> list[str]:  # type: ignore[type-arg]
    return [item["key"] for item in analysis["columns"]]


# ─── CSV: кодировки, разделители, локаль ─────────────────────────────────────

REPORT = """Сводка обследования районов;;;;;
Министерство по чрезвычайным ситуациям;;;;;
;;;;;
Район;Население, чел.;Площадь, км2;Дата обследования;Обследован;Ущерб, сомони
Вахдат;1\u00a0234\u00a0567;1 234,5;18.09.2025;да;12 500,75
Рудаки;450\u00a0000;2 891,25;01.02.2025;нет;0
Гиссар;много;1 040;13.01.2025;да;3 200,5
Турсунзаде;290 000;960,4;22.03.2025;да;—
Шахринав;110 500;590;05.04.2025;нет;150
Варзоб;80 200;1 700,75;30.04.2025;да;1 200
Файзабад;нет данных;850,1;15.05.2025;нет;75,25
Рогун;46 000;1 200;07.06.2025;да;980
Нурек;60 000;750,5;19.07.2025;да;15 000
Яван;190 000;900;28.08.2025;нет;430,6
Дангара;130 000;1 100;09.09.2025;да;220
Куляб;210 000;830,4;11.10.2025;нет;1 050,9
"""


def test_cp1251_с_точкой_с_запятой_десятичной_запятой_и_строками_над_таблицей(
    tmp_path: Path,
) -> None:
    source = write(tmp_path / "svodka.csv", REPORT, "cp1251")
    analysis = analyze(source)
    assert analysis["format"] == "csv"
    assert analysis["encoding"] == "cp1251"
    assert analysis["delimiter"] == ";"
    assert analysis["decimal"] == ","
    assert analysis["thousands"] == " "
    assert analysis["dateOrder"] == "dmy"
    assert (analysis["skipRows"], analysis["headerRows"]) == (3, 1)
    assert (analysis["rowEstimate"], analysis["approx"]) == (12, False)
    assert keys(analysis) == [
        "rayon",
        "naselenie_chel",
        "ploshchad_km2",
        "data_obsledovaniya",
        "obsledovan",
        "ushcherb_somoni",
    ]
    types = {item["key"]: (item["type"], item["semantic"]) for item in analysis["columns"]}
    assert types == {
        "rayon": ("text", "dimension"),
        "naselenie_chel": ("integer", "measure"),
        "ploshchad_km2": ("number", "measure"),
        "data_obsledovaniya": ("date", "time"),
        "obsledovan": ("boolean", "category"),
        "ushcherb_somoni": ("number", "measure"),
    }
    population = column(analysis, "Население, чел.")
    assert population["invalid"] == 1
    assert population["emptyShare"] == pytest.approx(1 / 12, abs=1e-3)
    assert population["samples"][:2] == ["1\u00a0234\u00a0567", "450\u00a0000"]
    assert column(analysis, "Дата обследования")["format"] == {"dateFormat": "dd.MM.yyyy"}
    assert column(analysis, "Район")["unique"] is True
    assert any("пропущены: 3" in warning for warning in analysis["warnings"])
    assert any("«Население, чел.»: 1 значение" in warning for warning in analysis["warnings"])
    assert analysis["preview"][0][0] == "Сводка обследования районов"
    assert analysis["preview"][3][0] == "Район"
    assert len(analysis["preview"]) == 16

    done = normalize(source, mapping_from(analysis))
    assert (done.result.rows, done.result.errors, done.result.written) == (12, 1, 11)
    # Номер строки — как в файле: заголовок в 4-й строке, «Гиссар» — в 7-й
    assert done.errors == [
        {"row": "7", "field": "naselenie_chel", "value": "много", "code": "invalid_integer"}
    ]
    assert done.result.error_sample == [
        {"row": 7, "column": "naselenie_chel", "value": "много", "reason": "invalid_integer"}
    ]
    assert done.rows[0] == ["5", "Вахдат", "1234567", "1234.5", "2025-09-18", "true", "12500.75"]
    assert done.rows[1] == ["6", "Рудаки", "450000", "2891.25", "2025-02-01", "false", "0"]
    turs = next(row for row in done.rows if row[1] == "Турсунзаде")
    assert turs[6] is None  # «—» в числе — NULL, а не ошибка
    faiz = next(row for row in done.rows if row[1] == "Файзабад")
    assert faiz[2] is None
    assert [row[0] for row in done.rows][-1] == "16"


def test_параметры_анализа_дают_ту_же_нормализацию(tmp_path: Path) -> None:
    source = write(tmp_path / "svodka.csv", REPORT, "cp1251")
    analysis = analyze(source)
    options = {
        name: analysis[name]
        for name in ("format", "encoding", "delimiter", "skipRows", "headerRows", "decimal")
    }
    options.update(thousands=analysis["thousands"], dateOrder=analysis["dateOrder"])
    auto = normalize(source, mapping_from(analysis))
    pinned = normalize(source, mapping_from(analysis), options=options)
    assert pinned.text == auto.text
    assert pinned.errors == auto.errors


def test_utf8_с_bom_запятая_и_даты_iso(tmp_path: Path) -> None:
    source = write(
        tmp_path / "bom.csv",
        '\ufeffid,Название,Сумма,Дата\n1,Альфа,"1,234.50",2026-09-18\n'
        "2,Бета,99.9,2026-09-19\n3,Гамма,100,2026-09-20\n",
    )
    analysis = analyze(source)
    assert analysis["encoding"] == "utf-8-sig"
    assert analysis["delimiter"] == ","
    assert (analysis["decimal"], analysis["thousands"]) == (".", ",")
    assert analysis["dateOrder"] == "ymd"
    assert keys(analysis) == ["id", "nazvanie", "summa", "data"]
    assert column(analysis, "id")["type"] == "integer"
    assert column(analysis, "id")["semantic"] == "identifier"
    assert column(analysis, "Сумма")["format"] == {"precision": 2, "thousands": True}
    done = normalize(source, mapping_from(analysis))
    assert done.rows[0] == ["2", "1", "Альфа", "1234.50", "2026-09-18"]


def test_tsv_с_таджикскими_заголовками(tmp_path: Path) -> None:
    source = write(
        tmp_path / "nohiyaho.tsv",
        "Ноҳия\tШумораи аҳолӣ\tМасоҳат\nВаҳдат\t1234567\t1234.5\nРӯдакӣ\t450000\t2891.25\n"
        "Ҳисор\t290000\t960.4\n",
    )
    analysis = analyze(source)
    assert (analysis["format"], analysis["delimiter"]) == ("tsv", "\t")
    assert keys(analysis) == ["nohiya", "shumorai_aholi", "masohat"]
    done = normalize(source, mapping_from(analysis))
    assert done.rows[1] == ["3", "Рӯдакӣ", "450000", "2891.25"]


def test_utf16_из_excel_юникод_текст(tmp_path: Path) -> None:
    text = "Код\tНаименование\tКоличество\r\n" + "".join(
        f"{index:03d}\tПозиция {index}\t{index * 10}\r\n" for index in range(1, 8)
    )
    source = write(tmp_path / "export.txt", text, "utf-16")
    analysis = analyze(source)
    assert analysis["encoding"] == "utf-16"
    assert analysis["delimiter"] == "\t"
    # Коды с ведущими нулями — идентификаторы, а не числа
    assert column(analysis, "Код")["type"] == "identifier"
    assert column(analysis, "Количество")["type"] == "integer"
    done = normalize(source, mapping_from(analysis))
    assert done.rows[0] == ["2", "001", "Позиция 1", "10"]
    assert done.result.rows == 7


def test_переводы_строк_и_кавычки_внутри_ячеек(tmp_path: Path) -> None:
    source = write(
        tmp_path / "notes.csv",
        "Код,Описание,Сумма\n"
        '1,"Первая строка\nвторая строка",100\n'
        '2,"Текст с ""кавычками"", и запятой",200\n'
        '3,"Ещё\r\nодна",abc\n'
        "4,Простой текст,400\n",
    )
    analysis = analyze(source)
    assert column(analysis, "Описание")["type"] == "long_text"
    assert analysis["rowEstimate"] == 4
    mapping = [
        {"column": 0, "fieldKey": "kod", "label": {"ru": "Код"}, "type": "integer"},
        {"column": 1, "fieldKey": "opisanie", "label": {"ru": "Описание"}, "type": "long_text"},
        {"column": 2, "fieldKey": "summa", "label": {"ru": "Сумма"}, "type": "integer"},
    ]
    for item in mapping:
        item["semantic"] = "dimension"
    done = normalize(source, mapping)
    # Номер строки — номер записи, как строки в Excel (многострочная ячейка — одна строка)
    assert done.rows == [
        ["2", "1", "Первая строка\nвторая строка", "100"],
        ["3", "2", 'Текст с "кавычками", и запятой', "200"],
        ["5", "4", "Простой текст", "400"],
    ]
    assert done.errors == [
        {"row": "4", "field": "summa", "value": "abc", "code": "invalid_integer"}
    ]


def test_даты_в_разных_форматах_и_серийные_числа_excel(tmp_path: Path) -> None:
    days = [date(2025, 9, 18) + timedelta(days=17 * index) for index in range(10)]
    lines = ["№;Дата события;Дата отчёта;Срок;Дата регистрации;Разные даты"]
    for index, day in enumerate(days, start=1):
        mixed = day.strftime("%d.%m.%Y") if index % 2 else day.isoformat()
        serial = (day - EXCEL_EPOCH).days
        lines.append(f"{index};{day:%d.%m.%Y};{day.isoformat()};{day:%d/%m/%y};{serial};{mixed}")
    source = write(tmp_path / "dates.csv", "\n".join(lines) + "\n")
    analysis = analyze(source)
    assert analysis["dateOrder"] == "dmy"
    formats = {item["name"]: (item["type"], item.get("format")) for item in analysis["columns"]}
    assert formats["Дата события"] == ("date", {"dateFormat": "dd.MM.yyyy"})
    assert formats["Дата отчёта"] == ("date", {"dateFormat": "yyyy-MM-dd"})
    assert formats["Срок"] == ("date", {"dateFormat": "dd/MM/yy"})
    assert formats["Дата регистрации"][0] == "date"
    assert formats["Разные даты"][0] == "date"
    assert column(analysis, "№")["semantic"] == "identifier"
    assert any("даты Excel" in warning for warning in analysis["warnings"])
    assert any("разных форматах" in warning for warning in analysis["warnings"])
    done = normalize(source, mapping_from(analysis))
    assert done.result.errors == 0
    for row, day in zip(done.rows, days, strict=True):
        assert row[2:] == [day.isoformat()] * 5


def test_порядок_месяц_день_по_данным(tmp_path: Path) -> None:
    source = write(
        tmp_path / "us.csv", "Date,Amount\n09/18/2025,10\n12/25/2025,20\n01/02/2025,30\n"
    )
    analysis = analyze(source)
    assert analysis["dateOrder"] == "mdy"
    assert column(analysis, "Date")["format"] == {"dateFormat": "MM/dd/yyyy"}
    done = normalize(source, mapping_from(analysis))
    assert [row[1] for row in done.rows] == ["2025-09-18", "2025-12-25", "2025-01-02"]


def test_числа_с_неразрывными_пробелами_процентами_и_скобками(tmp_path: Path) -> None:
    source = write(
        tmp_path / "numbers.csv",
        "Показатель;Значение;Доля;Баланс\n"
        "А;1\u00a0234\u00a0567,5;12,5%;(1 500)\n"
        "Б;2\u202f000,25;7%;−250\n"
        "В;1'000,75;100%;3 000\n"
        "Г;15;0,5%;0\n",
    )
    analysis = analyze(source)
    assert column(analysis, "Значение")["type"] == "number"
    assert column(analysis, "Доля")["type"] == "percent"
    assert column(analysis, "Доля")["format"]["scale"] == "percent"
    assert column(analysis, "Баланс")["type"] == "integer"
    done = normalize(source, mapping_from(analysis))
    assert [row[2:] for row in done.rows] == [
        ["1234567.5", "12.5", "-1500"],
        ["2000.25", "7", "-250"],
        ["1000.75", "100", "3000"],
        ["15", "0.5", "0"],
    ]
    as_fraction = mapping_from(analysis, {"dolya": {"format": {"scale": "fraction"}}})
    shares = [row[3] for row in normalize(source, as_fraction).rows]
    assert shares == ["0.125", "0.07", "1", "0.005"]


def test_да_нет_и_ноль_один(tmp_path: Path) -> None:
    source = write(
        tmp_path / "flags.csv",
        "Название;Активен;Флаг;Признак;Количество\n"
        "а;да;true;1;0\nб;нет;false;0;1\nв;Да;TRUE;1;1\nг;НЕТ;False;0;0\n",
    )
    analysis = analyze(source)
    types = {item["name"]: item["type"] for item in analysis["columns"]}
    assert types == {
        "Название": "text",
        "Активен": "boolean",
        "Флаг": "boolean",
        "Признак": "boolean",
        "Количество": "integer",
    }
    done = normalize(source, mapping_from(analysis))
    assert done.rows[3] == ["5", "г", "false", "false", "false", "0"]


def test_пустые_строки_не_считаются_а_номера_строк_сохраняются(tmp_path: Path) -> None:
    source = write(tmp_path / "gaps.csv", "Район;Число\nА;1\n\nБ;2\n;\nВ;x\nГ;3\n\n\n")
    analysis = analyze(source)
    assert (analysis["rowEstimate"], analysis["approx"]) == (4, False)
    done = normalize(source, mapping_from(analysis, {"chislo": {"type": "integer"}}))
    assert (done.result.rows, done.result.errors) == (4, 1)
    assert [row[0] for row in done.rows] == ["2", "4", "7"]
    assert done.errors[0]["row"] == "6"


def test_без_заголовка(tmp_path: Path) -> None:
    source = write(tmp_path / "raw.csv", "1;2026-09-18;15,5\n2;2026-09-19;16,25\n3;2026-09-20;7\n")
    analysis = analyze(source)
    assert analysis["headerRows"] == 0
    assert [item["name"] for item in analysis["columns"]] == ["Столбец 1", "Столбец 2", "Столбец 3"]
    assert keys(analysis) == ["column_1", "column_2", "column_3"]
    assert any("не найдена" in warning for warning in analysis["warnings"])
    done = normalize(source, mapping_from(analysis))
    assert done.rows[0] == ["1", "1", "2026-09-18", "15.5"]


def test_заголовок_из_годов(tmp_path: Path) -> None:
    source = write(
        tmp_path / "years.csv",
        "Район;2023;2024;2025\nВахдат;1500,5;1600,25;1700\nРудаки;900;950,5;1000\n"
        "Гиссар;400;420;450,75\n",
    )
    analysis = analyze(source)
    assert analysis["headerRows"] == 1
    assert keys(analysis) == ["rayon", "f_2023", "f_2024", "f_2025"]


def test_коды_телефоны_счета_и_почта(tmp_path: Path) -> None:
    lines = ["Код района;Телефон;Номер счёта;Почта;Сайт;ИНН"]
    for index in range(12):
        lines.append(
            f"{index % 3 + 1};+992 93 500-{index:02d}-{index:02d};99293500{index:04d};"
            f"user{index}@mchs.tj;https://mchs.tj/{index};{100000000 + index}"
        )
    source = write(tmp_path / "contacts.csv", "\n".join(lines) + "\n")
    analysis = analyze(source)
    types = {item["name"]: (item["type"], item["semantic"]) for item in analysis["columns"]}
    assert types == {
        "Код района": ("integer", "category"),
        "Телефон": ("phone", "identifier"),
        "Номер счёта": ("identifier", "identifier"),
        "Почта": ("email", "identifier"),
        "Сайт": ("url", "identifier"),
        "ИНН": ("identifier", "identifier"),
    }


def test_повторы_и_необычные_названия_столбцов(tmp_path: Path) -> None:
    source = write(
        tmp_path / "names.csv",
        "Сумма;Сумма;2025 год;№ п/п;_id;;Name;\n1;2;3;1;10;x;a;\n2;3;4;2;11;y;b;\n",
    )
    analysis = analyze(source)
    assert keys(analysis) == ["summa", "summa_2", "f_2025_god", "no_p_p", "id", "column_6", "name"]
    assert column(analysis, "Столбец 6")["index"] == 5
    assert any("«Сумма»" in warning for warning in analysis["warnings"])


def test_utf8_с_хвостом_в_windows_1251(tmp_path: Path) -> None:
    # Начало файла — чистый ASCII (кодировка по нему — UTF-8), в конце — строки из cp1251
    head = "name,city,value\n" + "Alpha,Dushanbe,1\n" * 70_000
    source = tmp_path / "mixed.csv"
    source.write_bytes(head.encode() + "Бета,Худжанд,5\n".encode("cp1251"))
    analysis = analyze(source)
    assert analysis["encoding"] == "utf-8"
    done = normalize(source, mapping_from(analysis))
    assert done.rows[-1] == [str(70_002), "Бета", "Худжанд", "5"]


# ─── Excel ───────────────────────────────────────────────────────────────────


def test_книга_из_трёх_листов_со_строками_над_заголовком(tmp_path: Path) -> None:
    book = Workbook()
    cover = book.active
    assert cover is not None
    cover.title = "Титул"
    cover["A1"] = "Сведения о происшествиях"
    data = book.create_sheet("Данные")
    data.append(["Сведения о ЧС за 2025 год"])
    data.append(["Составлено: отдел мониторинга"])
    data.append([])
    data.append(["Дата", "Район", "Пострадавших", "Ущерб, сомони", "Подтверждено", "Время"])
    districts = ["Вахдат", "Рудаки", "Гиссар", "Варзоб", "Рогун", "Нурек", "Яван", "Куляб"]
    districts += ["Дангара", "Фархор", "Шахритус", "Бальджуван"]
    for index, name in enumerate(districts):
        injured: object = index * 2
        if name == "Рогун":
            injured = "много"
        elif name == "Нурек":
            injured = "н/д"
        data.append(
            [
                datetime(2025, 1, 5) + timedelta(days=9 * index),
                name,
                injured,
                1500.5 + index * 100,
                index % 3 != 0,
                time(8 + index, 30),
            ]
        )
    reference = book.create_sheet("Справочник")
    reference.append(["Код", "Название"])
    for code, title in [("01", "Душанбе"), ("02", "Хатлон"), ("03", "Согд")]:
        reference.append([code, title])
    source = tmp_path / "report.xlsx"
    book.save(source)

    analysis = analyze(source)
    assert analysis["format"] == "xlsx"
    assert analysis["encoding"] is None and analysis["delimiter"] is None
    assert analysis["sheets"] == [
        {"name": "Титул", "rows": 1},
        {"name": "Данные", "rows": 16},
        {"name": "Справочник", "rows": 4},
    ]
    assert analysis["sheet"] == "Данные"
    assert (analysis["skipRows"], analysis["headerRows"]) == (3, 1)
    assert (analysis["rowEstimate"], analysis["approx"]) == (12, False)
    types = {item["name"]: item["type"] for item in analysis["columns"]}
    assert types == {
        "Дата": "date",
        "Район": "text",
        "Пострадавших": "integer",
        "Ущерб, сомони": "number",
        "Подтверждено": "boolean",
        "Время": "time",
    }
    assert column(analysis, "Пострадавших")["invalid"] == 1

    done = normalize(source, mapping_from(analysis))
    # Номера строк — строки листа: заголовок в 4-й, «Рогун» — в 9-й
    assert done.errors == [
        {"row": "9", "field": "postradavshikh", "value": "много", "code": "invalid_integer"}
    ]
    assert done.rows[0] == ["5", "2025-01-05", "Вахдат", "0", "1500.5", "false", "08:30:00"]
    nurek = next(row for row in done.rows if row[2] == "Нурек")
    assert nurek[3] is None

    listing = analyze(source, {"sheet": "Справочник"})
    assert listing["sheet"] == "Справочник"
    assert column(listing, "Код")["type"] == "identifier"
    with pytest.raises(ImportFileError) as error:
        analyze(source, {"sheet": "Нет такого"})
    assert error.value.code == "sheet_not_found"


def test_двухстрочный_заголовок_с_объединёнными_ячейками(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "Район"
    sheet.merge_cells("A1:A2")
    sheet["B1"] = "Население"
    sheet.merge_cells("B1:C1")
    sheet["B2"], sheet["C2"] = "Мужчины", "Женщины"
    sheet["D1"], sheet["D2"] = "Площадь", "км2"
    for index, name in enumerate(["Вахдат", "Рудаки", "Гиссар", "Варзоб", "Рогун"]):
        sheet.append([name, 1000 + index, 1100 + index, 200.5 + index])
    source = tmp_path / "header2.xlsx"
    book.save(source)
    analysis = analyze(source)
    assert analysis["headerRows"] == 2
    assert [item["name"] for item in analysis["columns"]] == [
        "Район",
        "Население / Мужчины",
        "Население / Женщины",
        "Площадь / км2",
    ]
    assert keys(analysis) == [
        "rayon",
        "naselenie_muzhchiny",
        "naselenie_zhenshchiny",
        "ploshchad_km2",
    ]
    done = normalize(source, mapping_from(analysis))
    assert done.rows[0] == ["3", "Вахдат", "1000", "1100", "200.5"]


def test_книга_excel_97_2003(tmp_path: Path) -> None:
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    sheet = book.add_sheet("Лист1")
    date_style = xlwt.easyxf(num_format_str="DD.MM.YYYY")
    for col, title in enumerate(["Дата", "Сумма", "Район"]):
        sheet.write(0, col, title)
    rows = [
        (datetime(2025, 3, 1), 1500.5, "Вахдат"),
        (datetime(2025, 3, 15), 200.0, "Рудаки"),
        (datetime(2025, 4, 2), 75.25, "Гиссар"),
    ]
    for index, (day, amount, name) in enumerate(rows, start=1):
        sheet.write(index, 0, day, date_style)
        sheet.write(index, 1, amount)
        sheet.write(index, 2, name)
    source = tmp_path / "old.xls"
    book.save(str(source))
    analysis = analyze(source)
    assert analysis["format"] == "xls"
    assert analysis["sheets"] == [{"name": "Лист1", "rows": 4}]
    assert [item["type"] for item in analysis["columns"]] == ["date", "number", "text"]
    done = normalize(source, mapping_from(analysis))
    assert done.rows[1] == ["3", "2025-03-15", "200", "Рудаки"]


# ─── JSON, NDJSON, GeoJSON ───────────────────────────────────────────────────


def test_json_массив_объектов(tmp_path: Path) -> None:
    records = [
        {
            "id": index,
            "name": name,
            "population": population,
            "area": 1234.5 + index,
            "active": index % 2 == 0,
            "surveyed": f"2025-0{index}-1{index}",
            "meta": {"source": "МЧС", "n": index},
            "note": "" if index == 1 else None,
        }
        for index, (name, population) in enumerate(
            [("Вахдат", 1234567), ("Рудаки", None), ("Гиссар", 290000), ("Варзоб", 80200)],
            start=1,
        )
    ]
    source = write(tmp_path / "districts.json", json.dumps(records, ensure_ascii=False))
    analysis = analyze(source)
    assert analysis["format"] == "json"
    assert (analysis["skipRows"], analysis["headerRows"]) == (0, 0)
    types = {item["key"]: item["type"] for item in analysis["columns"]}
    assert types == {
        "id": "integer",
        "name": "text",
        "population": "integer",
        "area": "number",
        "active": "boolean",
        "surveyed": "date",
        "meta": "json",
        "note": "text",
    }
    assert analysis["preview"][0][:2] == ["1", "Вахдат"]
    done = normalize(source, mapping_from(analysis))
    first = done.rows[0]
    assert first[:7] == ["1", "1", "Вахдат", "1234567", "1235.5", "false", "2025-01-11"]
    assert json.loads(first[7] or "") == {"source": "МЧС", "n": 1}
    # Пустая строка JSON — пустая строка, null — NULL
    assert first[8] == ""
    assert done.rows[1][8] is None
    assert done.rows[1][3] is None


def test_json_с_обёрткой_data(tmp_path: Path) -> None:
    source = write(
        tmp_path / "wrapped.json",
        json.dumps({"total": 2, "data": [{"code": "A", "value": 1}, {"code": "B", "value": 2}]}),
    )
    analysis = analyze(source)
    assert keys(analysis) == ["code", "value"]
    assert normalize(source, mapping_from(analysis)).rows == [["1", "A", "1"], ["2", "B", "2"]]


def test_ndjson_с_битой_строкой(tmp_path: Path) -> None:
    source = write(
        tmp_path / "events.ndjson",
        '{"code": "A1", "value": 10}\n{"code": "A2", "value": 20}\n\n'
        '{"code": "A3", "value": oops}\n{"code": "A4", "value": "30"}\n',
    )
    analysis = analyze(source)
    assert analysis["format"] == "ndjson"
    assert any("не читаются как JSON: 4" in warning for warning in analysis["warnings"])
    mapping = mapping_from(analysis, {"value": {"type": "integer"}})
    done = normalize(source, mapping)
    assert (done.result.rows, done.result.errors) == (4, 1)
    assert done.errors[0]["row"] == "4"
    assert done.errors[0]["code"] == "invalid_json"
    assert done.rows == [["1", "A1", "10"], ["2", "A2", "20"], ["5", "A4", "30"]]


GEOJSON = {
    "type": "FeatureCollection",
    "features": [
        {
            "type": "Feature",
            "properties": {"name": "Пост 1", "level": 3},
            "geometry": {"type": "Point", "coordinates": [68.78, 38.56]},
        },
        {
            "type": "Feature",
            "properties": {"name": "Зона", "level": 1},
            "geometry": {
                "type": "Polygon",
                "coordinates": [[[68, 38], [69, 38], [69, 39], [68, 38]]],
            },
        },
        {
            "type": "Feature",
            "properties": {"name": "Разрыв", "level": 2},
            "geometry": {
                "type": "Polygon",
                "coordinates": [[[68, 38], [69, 38], [69, 39], [68, 39]]],
            },
        },
        {"type": "Feature", "properties": {"name": "Без места", "level": 0}, "geometry": None},
        {
            "type": "Feature",
            "properties": {"name": "Высота", "level": 5},
            "geometry": {"type": "Point", "coordinates": [68.7, 38.5, 850]},
        },
    ],
}


def test_geojson_featurecollection(tmp_path: Path) -> None:
    source = write(tmp_path / "posts.geojson", json.dumps(GEOJSON, ensure_ascii=False))
    analysis = analyze(source)
    assert analysis["format"] == "geojson"
    assert analysis["geometry"] == {"kind": "features"}
    assert keys(analysis) == ["name", "level"]
    assert any("Геометрия 1 объекта" in warning for warning in analysis["warnings"])
    done = normalize(
        source, mapping_from(analysis), geometry={"kind": "features"}, geometry_field="geom"
    )
    assert done.rows[0] == ["1", "Пост 1", "3", "SRID=4326;POINT(68.78 38.56)"]
    assert done.rows[1][3] == "SRID=4326;POLYGON((68 38,69 38,69 39,68 38))"
    assert done.rows[2] == ["4", "Без места", "0", None]
    assert done.rows[3][3] == "SRID=4326;POINT(68.7 38.5)"
    assert [(item["row"], item["field"], item["code"]) for item in done.errors] == [
        ("3", "geom", "invalid_geometry")
    ]


# ─── Геометрия в CSV ─────────────────────────────────────────────────────────


def test_широта_и_долгота(tmp_path: Path) -> None:
    source = write(
        tmp_path / "posts.csv",
        "Пост;Широта;Долгота;Высота, м\n"
        "Душанбе;38,5598;68,7870;800\n"
        "Хорог;37,4897;71,5530;2200\n"
        "Худжанд;40,2826;69,6221;300\n"
        "Куляб;37,9146;69,7845;600\n"
        "Бохтар;37,8364;68,7806;430\n"
        "Истаравшан;39,9108;69,0036;1000\n"
        "Пенджикент;39,4952;67,6093;900\n"
        "Рогун;38,6907;69,7765;1400\n"
        "Турсунзаде;38,5108;68,2303;750\n"
        "Гарм;39,0064;70,3650;1300\n"
        "Ошибка;95,1;68,1;100\n"
        "Пусто;;;50\n",
    )
    analysis = analyze(source)
    assert analysis["geometry"] == {"kind": "latlon", "lat": 1, "lon": 2}
    assert column(analysis, "Широта")["semantic"] == "dimension"
    done = normalize(
        source,
        mapping_from(analysis),
        geometry=analysis["geometry"],
        geometry_field="geom",
    )
    assert done.rows[0] == [
        "2",
        "Душанбе",
        "38.5598",
        "68.7870",
        "800",
        "SRID=4326;POINT(68.787 38.5598)",
    ]
    assert done.rows[-1] == ["13", "Пусто", None, None, "50", None]
    assert done.errors == [
        {"row": "12", "field": "geom", "value": "95,1; 68,1", "code": "invalid_geometry"}
    ]


def test_wkt_в_столбце(tmp_path: Path) -> None:
    source = write(
        tmp_path / "shapes.csv",
        "id,Название,Геометрия\n"
        "1,Пост,POINT(68.78 38.56)\n"
        '2,Дорога,"LINESTRING(68 38, 69 39)"\n'
        '3,Зона,"POLYGON((68 38, 69 38, 69 39, 68 38))"\n'
        "4,Высота,POINT Z (68 38 850)\n"
        '5,Округ,"MULTIPOLYGON(((68 38, 69 38, 69 39, 68 38)))"\n'
        "6,EWKT,SRID=4326;POINT(68 38)\n"
        '7,Незамкнут,"POLYGON((68 38, 69 38, 69 39, 68 39))"\n'
        "8,Меркатор,SRID=3857;POINT(1 2)\n"
        "9,Вне,POINT(200 10)\n"
        "10,Пост 2,POINT(69.1 38.2)\n",
    )
    analysis = analyze(source)
    shape = column(analysis, "Геометрия")
    assert (shape["type"], shape["semantic"], shape["invalid"]) == ("geometry", "geometry", 3)
    assert analysis["geometry"] == {"kind": "wkt", "column": 2}
    done = normalize(source, mapping_from(analysis))
    geometry = {row[1]: row[3] for row in done.rows}
    assert geometry["2"] == "SRID=4326;LINESTRING(68 38,69 39)"
    assert geometry["4"] == "SRID=4326;POINT(68 38)"
    assert geometry["6"] == "SRID=4326;POINT(68 38)"
    assert [(item["row"], item["code"]) for item in done.errors] == [
        ("8", "invalid_geometry"),
        ("9", "invalid_geometry"),
        ("10", "invalid_geometry"),
    ]


# ─── Ограничения полей ───────────────────────────────────────────────────────


def test_обязательные_длина_переполнение_время_и_пояс(tmp_path: Path) -> None:
    long_text = "я" * 10_001
    source = write(
        tmp_path / "limits.csv",
        "Код,Название,Количество,Время,Когда,Описание\n"
        "A1,Первый,9223372036854775807,14:30,2026-09-18 14:30,ok\n"
        "A2,,5,2:05 PM,2026-09-18T14:30:00+03:00,ok\n"
        f",Третий,9223372036854775808,25:00,18.09.2026 08:00,{long_text}\n"
        "A4,Четвёртый,н/д,07:15:30,нет,ok\n",
    )
    mapping = [
        {"column": 0, "fieldKey": "kod", "type": "identifier", "required": True},
        {"column": 1, "fieldKey": "nazvanie", "type": "text", "required": True},
        {"column": 2, "fieldKey": "kolichestvo", "type": "integer"},
        {"column": 3, "fieldKey": "vremya", "type": "time"},
        {"column": 4, "fieldKey": "kogda", "type": "datetime"},
        {"column": 5, "fieldKey": "opisanie", "type": "text"},
    ]
    for item in mapping:
        item.update(label={"ru": str(item["fieldKey"])}, semantic="dimension")
    done = normalize(source, mapping)
    assert done.rows == [
        ["2", "A1", "Первый", "9223372036854775807", "14:30:00", "2026-09-18T14:30:00+05:00", "ok"]
    ]
    assert [(item["row"], item["field"], item["code"]) for item in done.errors] == [
        ("3", "nazvanie", "required"),
        ("4", "kod", "required"),
        ("4", "kolichestvo", "invalid_integer"),
        ("4", "vremya", "invalid_time"),
        ("4", "opisanie", "too_long"),
        ("5", "kogda", "invalid_datetime"),
    ]
    assert (done.result.rows, done.result.errors, done.result.error_lines) == (4, 3, 6)
    # Длинное значение в файле ошибок обрезано
    assert len(done.errors[4]["value"]) == 500


def test_значение_для_excel_не_становится_формулой(tmp_path: Path) -> None:
    source = write(tmp_path / "formula.csv", "Число\n=HYPERLINK(1)\n-5x\n12\n")
    mapping = [{"column": 0, "fieldKey": "chislo", "label": {"ru": "Число"}, "type": "integer"}]
    mapping[0]["semantic"] = "measure"
    done = normalize(source, mapping)
    assert [item["value"] for item in done.errors] == ["'=HYPERLINK(1)", "'-5x"]
    assert done.result.error_sample[0]["value"] == "=HYPERLINK(1)"


# ─── Большие и испорченные файлы ─────────────────────────────────────────────


def test_большой_файл_анализируется_по_началу(tmp_path: Path) -> None:
    line = "12345;Вахдат;1 234,56;18.09.2025\n"
    count = SAMPLE_BYTES // len(line.encode()) + 50_000
    full = tmp_path / "big.csv"
    full.write_bytes(("Код;Район;Сумма;Дата\n" + line * count).encode())
    head = tmp_path / "head.csv"
    with full.open("rb") as stream:
        head.write_bytes(stream.read(SAMPLE_BYTES))
    analysis = analyze_file(head, "big.csv", {}, complete=False, file_size=full.stat().st_size)
    assert analysis["approx"] is True
    assert analysis["rowEstimate"] == pytest.approx(count, rel=0.01)
    assert len(analysis["preview"]) == 50
    assert column(analysis, "Сумма")["type"] == "number"


def test_испорченные_файлы(tmp_path: Path) -> None:
    empty = write(tmp_path / "empty.csv", "")
    with pytest.raises(ImportFileError) as error:
        analyze(empty)
    assert error.value.code == "empty"
    blank = write(tmp_path / "blank.csv", "\n\n  \n")
    with pytest.raises(ImportFileError) as error:
        analyze(blank)
    assert error.value.code == "empty"
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"PK\x03\x04" + b"\x00" * 100)
    with pytest.raises(ImportFileError) as error:
        analyze(broken)
    assert error.value.code == "unreadable"
    cut = write(tmp_path / "cut.json", '[{"a": 1}, {"a": 2')
    with pytest.raises(ImportFileError) as error:
        normalize(cut, [{"column": 0, "fieldKey": "a", "type": "integer"}])
    assert error.value.code == "unreadable"
