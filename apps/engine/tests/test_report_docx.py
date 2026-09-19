"""DOCX отчёта (ADR-0078): модель страницы печати → документ Word через docxtpl.

Текст Tiptap — абзацами, заголовками и списками; таблицы и показатели — таблицами;
графики и карты — картинками; колонтитулы — с полями номера страницы.
"""

import io

import pytest

pytest.importorskip("docxtpl")

from docx import Document
from PIL import Image

from kchs_engine.render.docx_report import (
    DocxLabels,
    add_rich_text,
    build_docx,
    content_width_mm,
)
from kchs_engine.render.report import (
    appearance_state,
    content_width_px,
    count_pdf_pages,
    footer_template,
    header_template,
)

LABELS = DocxLabels(header="Сводка ЧС", footer="Для служебного пользования", page="Стр.", of="из")


def png(width: int = 40, height: int = 20) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (width, height), (40, 90, 200)).save(out, format="PNG")
    return out.getvalue()


RICH = {
    "type": "doc",
    "content": [
        {
            "type": "heading",
            "attrs": {"level": 2},
            "content": [{"type": "text", "text": "Обстановка"}],
        },
        {
            "type": "paragraph",
            "content": [
                {"type": "text", "text": "За неделю "},
                {"type": "text", "text": "12 паводков", "marks": [{"type": "bold"}]},
                {"type": "hardBreak"},
                {
                    "type": "text",
                    "text": "портал",
                    "marks": [{"type": "link", "attrs": {"href": "https://kchs.tj"}}],
                },
            ],
        },
        {
            "type": "bulletList",
            "content": [
                {
                    "type": "listItem",
                    "content": [
                        {"type": "paragraph", "content": [{"type": "text", "text": "Хатлон"}]}
                    ],
                },
                {
                    "type": "listItem",
                    "content": [
                        {"type": "paragraph", "content": [{"type": "text", "text": "Согд"}]},
                        {
                            "type": "orderedList",
                            "content": [
                                {
                                    "type": "listItem",
                                    "content": [
                                        {
                                            "type": "paragraph",
                                            "content": [{"type": "text", "text": "Худжанд"}],
                                        }
                                    ],
                                }
                            ],
                        },
                    ],
                },
            ],
        },
        {
            "type": "blockquote",
            "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Цитата"}]}],
        },
        {"type": "codeBlock", "content": [{"type": "text", "text": "select 1"}]},
        {"type": "horizontalRule"},
        # Узел вне белого списка — пропускается
        {"type": "image", "attrs": {"src": "javascript:alert(1)"}},
    ],
}

MODEL = {
    "version": 1,
    "title": "Паводки <весна> & лето",
    "subtitle": "Период: апрель 2026 · Сформирован 19.09.2026",
    "settings": {
        "orientation": "portrait",
        "header": "",
        "footer": "",
        "titlePage": False,
        "formats": ["pdf", "docx"],
    },
    "labels": {"page": "Стр.", "of": "из"},
    "blocks": [
        {"id": "t1", "kind": "text", "body": RICH},
        {
            "id": "c1",
            "kind": "figure",
            "title": "Паводки по районам",
            "figure": "chart",
            "note": None,
        },
        {
            "id": "q1",
            "kind": "table",
            "title": "Таблица",
            "columns": [{"label": "Район", "numeric": False}, {"label": "Число", "numeric": True}],
            "rows": [["Хатлон", "12"], ["Согд", "7"]],
            "total": 5,
        },
        {
            "id": "m1",
            "kind": "metrics",
            "title": "Показатели",
            "items": [{"label": "Пострадавшие", "value": "34", "note": "+12 % к прошлому году"}],
        },
        {"id": "p1", "kind": "page_break"},
        {
            "id": "m2",
            "kind": "figure",
            "title": "Карта зон риска",
            "figure": "map",
            "note": "Слоёв: 2",
        },
        {"id": "x1", "kind": "notice", "title": "Бюджет", "text": "Нет доступа к данным"},
    ],
}


def test_build_docx_lays_out_blocks_and_page_fields() -> None:
    images = {"c1": png(), "m2": png(80, 60)}
    data = build_docx(MODEL, images.get, LABELS)
    doc = Document(io.BytesIO(data))
    texts = [paragraph.text for paragraph in doc.paragraphs]

    # Заголовок экранируется шаблоном, но в документе — как есть
    assert texts[0] == "Паводки <весна> & лето"
    assert "Период: апрель 2026" in texts[1]
    assert "Обстановка" in texts
    heading = next(p for p in doc.paragraphs if p.text == "Обстановка")
    assert heading.style.name == "Heading 2"
    assert any("12 паводков" in text for text in texts)
    assert any("(https://kchs.tj)" in text for text in texts)
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert styles["Хатлон"] == "List Bullet"
    assert styles["Худжанд"] == "List Number 2"
    assert styles["Цитата"] == "Quote"
    assert "select 1" in texts
    assert "Нет доступа к данным" in texts
    assert "javascript" not in "".join(texts)

    # Таблица результата и сетка показателей
    assert len(doc.tables) == 2
    table = doc.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Район", "Число"]
    assert [cell.text for cell in table.rows[2].cells] == ["Согд", "7"]
    assert "2 / 5" in texts
    assert doc.tables[1].rows[0].cells[1].text == "34"

    # Графики и карты — картинками, разрыв страницы — на месте
    assert len(doc.inline_shapes) == 2
    xml = doc.element.xml
    assert 'w:type="page"' in xml

    # Колонтитулы: текст и поля номера страницы
    footer, pages = doc.sections[0].footer.paragraphs[:2]
    assert footer.text == "Для служебного пользования"
    assert pages.text.startswith("Стр.")
    pages_xml = pages._p.xml
    assert "PAGE" in pages_xml and "NUMPAGES" in pages_xml
    assert doc.sections[0].header.paragraphs[0].text == "Сводка ЧС"


def test_build_docx_landscape_and_missing_images() -> None:
    model = {**MODEL, "settings": {**MODEL["settings"], "orientation": "landscape"}}
    doc = Document(io.BytesIO(build_docx(model, lambda _id: None, LABELS)))
    section = doc.sections[0]
    assert section.page_width > section.page_height
    # Нет картинки (не снялась) — подпись остаётся, картинки нет
    assert len(doc.inline_shapes) == 0
    assert "Паводки по районам" in [p.text for p in doc.paragraphs]
    assert content_width_mm("landscape") > content_width_mm("portrait")


def test_rich_text_ignores_foreign_nodes() -> None:
    doc = Document()
    add_rich_text(doc, {"type": "doc", "content": [{"type": "mystery", "content": []}, "x"]})
    add_rich_text(doc, {"type": "doc"})
    assert [p.text for p in doc.paragraphs] == []


def test_print_helpers() -> None:
    assert count_pdf_pages(b"%PDF /Type /Pages /Type /Page x /Type/Page y") == 2
    assert count_pdf_pages(b"no pages") is None
    assert "&lt;b&gt;" in header_template("<b>")
    footer = footer_template("Отдел", "Страница", "из")
    assert 'class="pageNumber"' in footer and 'class="totalPages"' in footer
    # Окно — ширина поля печати A4 (182 мм ≈ 688 px), альбомная — шире
    assert 680 <= content_width_px("portrait") <= 695
    assert content_width_px("landscape") > content_width_px("portrait")
    state = appearance_state("http://web:80/any", "tg")
    assert state["origins"][0]["origin"] == "http://web:80"
    assert '"locale": "tg"' in state["origins"][0]["localStorage"][0]["value"]
