"""Импорт пользователей из Excel (P0-E04 S04, ADR-0041): разбор XLSX и шаблон."""

from datetime import date, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from kchs_engine import users_import
from kchs_engine.contracts import users_import_contract
from kchs_engine.users_import import build_template, cell_text, normalize_header, parse_workbook


def workbook(path: Path, sheets: dict[str, list[list[Any]]]) -> Path:
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    book.save(path)
    return path


def test_cell_text_как_видит_человек() -> None:
    assert cell_text(None) is None
    assert cell_text("  ") is None
    assert cell_text(" Иванов ") == "Иванов"
    # Телефон, набранный числом: Excel хранит float — без «.0» и без экспоненты
    assert cell_text(992935001122.0) == "992935001122"
    assert cell_text(12345) == "12345"
    assert cell_text(2.5) == "2.5"
    assert cell_text(datetime(2026, 9, 18)) == "2026-09-18"
    assert cell_text(datetime(2026, 9, 18, 14, 30)) == "2026-09-18T14:30"
    assert cell_text(date(2026, 1, 2)) == "2026-01-02"
    assert cell_text(True) == "true"
    assert len(cell_text("я" * 5000) or "") == users_import.MAX_CELL_CHARS


def test_normalize_header() -> None:
    assert normalize_header("  ЛОГИН * ") == "логин"
    assert normalize_header("Электронная\nпочта:") == "электронная почта"
    assert normalize_header("Отчёство") == "отчество"


def test_разбор_кириллицы_пустых_строк_и_лишних_столбцов(tmp_path: Path) -> None:
    source = workbook(
        tmp_path / "users.xlsx",
        {
            "Сотрудники": [
                ["Список сотрудников КЧС"],
                [],
                [
                    "ЛОГИН*",
                    "фамилия",
                    "Имя",
                    "Отчёство",
                    "E-mail",
                    "Телефон",
                    "Подразделение",
                    "Роли",
                    "Табельный номер",
                ],
                [
                    "karimova.z",
                    "Каримова",
                    "Зарина",
                    "Алиевна",
                    "karimova@kchs.tj",
                    992935001122.0,
                    "ОДС",
                    "employee, data_steward",
                    1001,
                ],
                [None, None, None, None, None, None, None, None, None],
                [12345, "Раҳимов", "Ҷамшед", None, None, datetime(2026, 9, 18), None, None, None],
                ["   ", None, None],
            ]
        },
    )

    result = parse_workbook(source)

    assert result.file_error is None
    assert result.total_rows == 2
    assert result.columns["login"] == "ЛОГИН*"
    assert result.columns["email"] == "E-mail"
    assert result.warnings == [
        {"code": "unknown_column", "field": None, "params": {"column": "Табельный номер"}}
    ]
    first, second = result.rows
    assert first["row"] == 4
    assert first["values"]["lastName"] == "Каримова"
    assert first["values"]["phone"] == "992935001122"
    assert first["values"]["unit"] == "ОДС"
    assert first["values"]["roles"] == "employee, data_steward"
    assert "locale" not in first["values"]
    # Пустая строка пропущена, номер следующей — как на листе Excel
    assert second["row"] == 6
    assert second["values"]["login"] == "12345"
    assert second["values"]["firstName"] == "Ҷамшед"
    assert second["values"]["phone"] == "2026-09-18"


def test_нет_обязательных_столбцов(tmp_path: Path) -> None:
    source = workbook(tmp_path / "x.xlsx", {"Лист": [["Логин", "Телефон"], ["a.b", "1"]]})
    result = parse_workbook(source)
    assert result.file_error is not None
    assert result.file_error["code"] == "missing_columns"
    assert "Фамилия" in str(result.file_error["params"]["columns"])
    assert result.rows == []


def test_повтор_столбца(tmp_path: Path) -> None:
    source = workbook(
        tmp_path / "x.xlsx", {"Лист": [["Логин", "Фамилия", "Имя", "Login"], ["a", "b", "c", "d"]]}
    )
    result = parse_workbook(source)
    assert result.file_error == {
        "code": "duplicate_column",
        "field": None,
        "params": {"column": "Login"},
    }


def test_нет_заголовка(tmp_path: Path) -> None:
    source = workbook(tmp_path / "x.xlsx", {"Лист": [["что-то"], ["ещё"]]})
    result = parse_workbook(source)
    assert result.file_error is not None
    assert result.file_error["code"] == "no_header"


@pytest.mark.parametrize(
    "content",
    [b"not an excel file at all", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 512],
    ids=["мусор", "старый-xls"],
)
def test_нечитаемый_файл(tmp_path: Path, content: bytes) -> None:
    source = tmp_path / "broken.xlsx"
    source.write_bytes(content)
    result = parse_workbook(source)
    assert result.file_error is not None
    assert result.file_error["code"] == "unreadable"


def test_предел_строк(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    contract = dict(users_import_contract())
    contract["maxRows"] = 3
    monkeypatch.setattr(users_import, "users_import_contract", lambda: contract)
    rows: list[list[Any]] = [["Логин", "Фамилия", "Имя"]]
    rows += [[f"user{i}", "Ф", "И"] for i in range(4)]
    result = parse_workbook(workbook(tmp_path / "x.xlsx", {"Лист": rows}))
    assert result.file_error == {"code": "too_many_rows", "field": None, "params": {"max": 3}}
    assert result.rows == []


def test_лист_данных_выбирается_по_заголовку(tmp_path: Path) -> None:
    source = workbook(
        tmp_path / "x.xlsx",
        {
            "Справка": [["Столбец", "Как заполнять"], ["Логин", "латиница"]],
            "Пользователи": [["Логин", "Фамилия", "Имя"], ["a.b", "Б", "В"]],
        },
    )
    result = parse_workbook(source)
    assert result.file_error is None
    assert [row["values"]["login"] for row in result.rows] == ["a.b"]


def test_шаблон_читается_разбором_без_строк_и_замечаний(tmp_path: Path) -> None:
    target = build_template(
        tmp_path / "template.xlsx",
        roles=[{"key": "employee", "name": "Сотрудник"}],
        units=[{"code": "ODS", "name": "Оперативно-дежурная служба", "active": True}],
        positions=[{"name": "Специалист", "unitCode": "ODS"}],
    )
    result = parse_workbook(target)
    assert result.file_error is None
    assert result.warnings == []
    assert result.rows == []
    fields = [spec["key"] for spec in users_import_contract()["fields"]]
    assert set(result.columns) == set(fields)

    book = load_workbook(target)
    assert book.sheetnames == ["Пользователи", "Справка", "Роли", "Подразделения", "Должности"]
    headers = [cell.value for cell in book["Пользователи"][1]]
    assert headers[0] == "Логин *"
    assert "Подразделение (код)" in headers
    assert book["Роли"]["A2"].value == "employee"
    assert book["Подразделения"]["B2"].value == "Оперативно-дежурная служба"


def test_шаблон_требует_сервисный_токен(monkeypatch: pytest.MonkeyPatch) -> None:
    from fastapi.testclient import TestClient

    from kchs_engine.config import settings
    from kchs_engine.main import app

    monkeypatch.setenv("INTERNAL_SERVICE_TOKEN", "test-token-123")
    settings.cache_clear()
    try:
        client = TestClient(app)  # без with: воркеры BullMQ не запускаются
        denied = client.post("/templates/users-import", json={})
        assert denied.status_code == 401
        wrong = client.post(
            "/templates/users-import", json={}, headers={"x-kchs-service-token": "wrong"}
        )
        assert wrong.status_code == 401
        ok = client.post(
            "/templates/users-import",
            json={"roles": [{"key": "employee", "name": "Сотрудник"}]},
            headers={"x-kchs-service-token": "test-token-123"},
        )
        assert ok.status_code == 200
        assert ok.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert ok.content[:2] == b"PK"
    finally:
        settings.cache_clear()


def test_обработчик_зарегистрирован() -> None:
    from kchs_engine.jobs import JOB_HANDLERS, registered_queues
    from kchs_engine.jobs import users_import as _registered  # noqa: F401

    assert "imports:users.parse" in JOB_HANDLERS
    assert "imports" in registered_queues()
