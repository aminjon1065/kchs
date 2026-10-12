"""Рендеры модуля документов (ADR-0085): разбор и заполнение шаблонов DOCX в
песочнице Jinja, наложение штампа и водяного знака на PDF (pypdf).

Печать в Chromium — в образе движка: `pytest -m browser`."""

from __future__ import annotations

import asyncio
import io
import shutil
import subprocess
from pathlib import Path
from typing import Any

import pytest

pytest.importorskip("docxtpl")
pytest.importorskip("pypdf")

from docx import Document
from pypdf import PageObject, PdfReader, PdfWriter
from pypdf.generic import ContentStream, DictionaryObject, NameObject

from kchs_engine.render import documents
from kchs_engine.render.documents import (
    PageBox,
    RenderContentError,
    collect_placeholders,
    fill_template,
    find_text,
    image_fragment,
    inspect_template,
    merge_overlays,
    page_sizes,
    place_image,
    plan_images,
    size_key,
    target_pages,
    with_fragments,
)

# ─── Плейсхолдеры ────────────────────────────────────────────────────────────


def test_placeholders_are_longest_attribute_chains() -> None:
    source = (
        "{{ doc.subject }} {{ doc.fields.addressee }} {{ author.position|upper }}"
        " {{ doc.reg_number or '____' }} {{ doc['summary'] }} {{ today }}"
    )
    assert collect_placeholders(source) == [
        "author.position",
        "doc.fields.addressee",
        "doc.reg_number",
        "doc.subject",
        "doc.summary",
        "today",
    ]


def test_loop_variables_are_not_context() -> None:
    source = (
        "{% for item in doc.attachments %}{{ loop.index }}. {{ item.name }}{% endfor %}"
        "{% set signer_name = signer.short_name %}{{ signer_name }}"
    )
    assert collect_placeholders(source) == ["doc.attachments", "signer.short_name"]


def _docx(path: Path, body: list[str], header: str | None = None) -> Path:
    document = Document()
    for text in body:
        document.add_paragraph(text)
    if header is not None:
        document.sections[0].header.paragraphs[0].text = header
    document.save(str(path))
    return path


def test_inspect_reads_body_and_header(tmp_path: Path) -> None:
    source = _docx(
        tmp_path / "t.docx",
        ["{{ doc.subject }}", "{%p for a in doc.attachments %}", "{{ a.name }}", "{%p endfor %}"],
        header="{{ org.name }}",
    )
    assert inspect_template(source) == ["doc.attachments", "doc.subject", "org.name"]


def test_inspect_syntax_error_is_content_error(tmp_path: Path) -> None:
    source = _docx(tmp_path / "t.docx", ["{{ doc.subject "])
    with pytest.raises(RenderContentError, match="Ошибка в шаблоне"):
        inspect_template(source)


def test_inspect_not_docx(tmp_path: Path) -> None:
    source = tmp_path / "t.docx"
    source.write_bytes(b"not a zip")
    with pytest.raises(RenderContentError, match="Word"):
        inspect_template(source)


# ─── Заполнение ──────────────────────────────────────────────────────────────


def _text(path: Path) -> str:
    return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)


def test_fill_escapes_values_for_word_xml(tmp_path: Path) -> None:
    source = _docx(
        tmp_path / "t.docx",
        ["Тема: {{ doc.subject }}", "{{ doc.fields.addressee }}", "[{{ unknown.path }}]"],
    )
    context: dict[str, Any] = {
        "doc": {"subject": "Смета <2026> & «план»", "fields": {"addressee": "Хукумат"}},
    }
    result = fill_template(source, context, tmp_path / "out.docx")
    text = _text(result)
    assert "Тема: Смета <2026> & «план»" in text
    assert "Хукумат" in text
    # Неизвестный плейсхолдер — пусто, не ошибка
    assert "[]" in text


def test_fill_loops_over_lists(tmp_path: Path) -> None:
    source = _docx(
        tmp_path / "t.docx",
        [
            "{%p for item in doc.attachments %}",
            "{{ loop.index }}. {{ item.name }}",
            "{%p endfor %}",
        ],
    )
    context = {"doc": {"attachments": [{"name": "Смета.xlsx"}, {"name": "Схема.pdf"}]}}
    text = _text(fill_template(source, context, tmp_path / "out.docx"))
    assert "1. Смета.xlsx" in text
    assert "2. Схема.pdf" in text


def test_fill_sandbox_blocks_python_internals(tmp_path: Path) -> None:
    source = _docx(tmp_path / "t.docx", ["{{ ''.__class__.__mro__[1].__subclasses__() }}"])
    with pytest.raises(RenderContentError, match="Ошибка в шаблоне"):
        fill_template(source, {}, tmp_path / "out.docx")


# ─── Наложение ───────────────────────────────────────────────────────────────


def _page_with(width: float, height: float, operators: bytes) -> PageObject:
    """Лист заданного размера с содержимым — операторами PDF."""
    writer = PdfWriter()
    page = writer.add_blank_page(width, height)
    stream = ContentStream(None, writer)
    stream.set_data(operators)
    page.replace_contents(stream)
    return page


def _pdf(pages: list[PageObject], rotate: dict[int, int] | None = None) -> bytes:
    writer = PdfWriter()
    for index, page in enumerate(pages):
        added = writer.add_page(page)
        if rotate and index in rotate:
            added.rotate(rotate[index])
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_page_sizes_follow_rotation(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    a4 = (595.0, 842.0)
    source.write_bytes(
        _pdf(
            [_page_with(*a4, b"0 0 1 rg 0 0 10 10 re f"), _page_with(*a4, b"")],
            rotate={1: 90},
        )
    )
    assert page_sizes(source, "first") == [a4]
    # Повёрнутый лист — альбомный: наложение печатается под видимый размер
    assert page_sizes(source, "all") == [a4, (842.0, 595.0)]


def test_merge_overlays_puts_overlay_on_target_pages(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    source.write_bytes(
        _pdf([_page_with(595, 842, b"0 0 1 rg 0 0 10 10 re f"), _page_with(595, 842, b"")])
    )
    overlay = _pdf([_page_with(595, 842, b"1 0 0 rg 100 100 50 50 re f")])
    overlays = {size_key((595.0, 842.0)): overlay}

    first, pages = merge_overlays(source, overlays, "first", tmp_path / "first.pdf")
    assert pages == 2
    reader = PdfReader(str(first))
    assert b"100 100 50 50 re" in reader.pages[0].get_contents().get_data()
    second = reader.pages[1].get_contents()
    assert second is None or b"100 100 50 50 re" not in second.get_data()

    every, _ = merge_overlays(source, overlays, "all", tmp_path / "all.pdf")
    for page in PdfReader(str(every)).pages:
        assert b"100 100 50 50 re" in page.get_contents().get_data()
    # Исходник не тронут
    assert b"100 100 50 50 re" not in source.read_bytes()


def test_merge_scales_overlay_to_page(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    source.write_bytes(_pdf([_page_with(600, 800, b"")]))
    # Chromium печатает с точностью до пикселя: наложение чуть другого размера
    overlays = {size_key((600.0, 800.0)): _pdf([_page_with(599.25, 799.5, b"0 g 0 0 5 5 re f")])}
    result, _ = merge_overlays(source, overlays, "first", tmp_path / "out.pdf")
    content = PdfReader(str(result)).pages[0].get_contents().get_data()
    assert b"cm" in content


def test_last_page_target() -> None:
    assert target_pages(3, "last") == [2]
    assert target_pages(3, "first") == [0]
    assert target_pages(3, "all") == [0, 1, 2]
    assert target_pages(0, "last") == []


# ─── Факсимиле: где поставить (ADR-0198) ─────────────────────────────────────

PNG = "data:image/png;base64,iVBORw0KGgo="
A4 = PageBox(0.0, 0.0, 595.0, 842.0)
text_layer = pytest.mark.skipif(shutil.which("pdftotext") is None, reason="нет poppler")


def _text_page(width: float, height: float, lines: list[tuple[float, float, str]]) -> PageObject:
    """Лист с текстом стандартным шрифтом: (x, базовая линия, текст)."""
    writer = PdfWriter()
    page = writer.add_blank_page(width, height)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    ops = b"".join(f"BT /F1 12 Tf {x} {y} Td ({text}) Tj ET\n".encode() for x, y, text in lines)
    stream = ContentStream(None, writer)
    stream.set_data(ops)
    page.replace_contents(stream)
    return page


def _image(**extra: Any) -> dict[str, Any]:
    return {
        "id": "sig-1",
        "dataUri": PNG,
        "widthMm": 40.0,
        "heightMm": 16.0,
        "anchors": [
            {"text": "[[SIGN]]", "mode": "cover"},
            {"text": "Saidzoda T. N.", "mode": "left"},
        ],
        "manual": None,
        "fallback": {"x": 0.7, "y": 0.75},
        **extra,
    }


def _words(*pages: list[tuple[str, float, float, float, float]]) -> list[list[Any]]:
    return [list(page) for page in pages]


def test_find_text_matches_words_of_one_line() -> None:
    pages = _words(
        [("Dear", 0.1, 0.1, 0.15, 0.12)],
        [
            ("Saidzoda", 0.6, 0.7, 0.72, 0.72),
            ("T.", 0.73, 0.7, 0.75, 0.72),
            ("N.", 0.76, 0.7, 0.78, 0.72),
            ("Saidzoda", 0.1, 0.3, 0.2, 0.32),
        ],
    )
    hits = find_text(pages, "saidzoda  t.\u00a0n.")
    assert len(hits) == 1
    assert (hits[0].page, hits[0].left, hits[0].right) == (1, 0.6, 0.78)


def test_marker_wins_and_is_covered() -> None:
    pages = _words([("[[SIGN]]", 0.4, 0.6, 0.5, 0.62), ("Saidzoda", 0.6, 0.6, 0.7, 0.62)])
    placement = place_image(pages, [A4], _image())
    assert (placement.source, placement.page) == ("marker", 0)
    assert placement.cover == (0.4, 0.6, 0.5, 0.62)
    assert placement.x == pytest.approx(0.45)


def test_facsimile_goes_left_of_the_name_on_the_latest_page() -> None:
    name = [
        ("Saidzoda", 0.65, 0.7, 0.75, 0.72),
        ("T.", 0.76, 0.7, 0.78, 0.72),
        ("N.", 0.79, 0.7, 0.8, 0.72),
    ]
    pages = _words(name, name)
    placement = place_image(
        pages, [A4, A4], _image(anchors=[{"text": "Saidzoda T. N.", "mode": "left"}])
    )
    assert (placement.source, placement.page) == ("name", 1)
    # Правый край факсимиле — левее начала фамилии
    assert placement.x + 40 * documents.MM / 595 / 2 < 0.65
    assert placement.y == pytest.approx(0.71)


def test_no_room_on_the_left_puts_facsimile_above() -> None:
    pages = _words(
        [
            ("Saidzoda", 0.05, 0.7, 0.15, 0.72),
            ("T.", 0.16, 0.7, 0.18, 0.72),
            ("N.", 0.19, 0.7, 0.2, 0.72),
        ]
    )
    placement = place_image(
        pages, [A4], _image(anchors=[{"text": "Saidzoda T. N.", "mode": "left"}])
    )
    assert placement.source == "name"
    assert placement.y < 0.7


def test_fallback_and_manual() -> None:
    pages = _words([("nothing", 0.1, 0.1, 0.2, 0.12)], [])
    fallback = place_image(pages, [A4, A4], _image())
    assert (fallback.source, fallback.page, fallback.x, fallback.y) == ("fallback", 1, 0.7, 0.75)
    manual = place_image(pages, [A4, A4], _image(manual={"page": 1, "x": 0.5, "y": 0.25}))
    assert (manual.source, manual.page, manual.x, manual.y) == ("manual", 0, 0.5, 0.25)
    # Лист за пределами документа — последний
    beyond = place_image(pages, [A4, A4], _image(manual={"page": 9, "x": 0.5, "y": 0.25}))
    assert beyond.page == 1


@text_layer
def test_plan_images_finds_the_name_in_a_real_pdf(tmp_path: Path) -> None:
    source = tmp_path / "signed.pdf"
    source.write_bytes(
        _pdf(
            [
                _text_page(595, 842, [(72, 700, "text")]),
                _text_page(595, 842, [(400, 200, "Saidzoda T. N.")]),
            ]
        )
    )
    count, sizes, fragments, placements = plan_images(source, "last", [_image()])
    assert count == 2
    assert list(sizes) == [1]
    assert placements[0]["page"] == 2 and placements[0]["source"] == "name"
    assert 0 < placements[0]["x"] < 400 / 595
    # Базовая линия 200 pt от низа — около 0,76 высоты от верха
    assert placements[0]["y"] == pytest.approx(1 - 200 / 842, abs=0.02)
    assert PNG in fragments[1][0]
    page = with_fragments("<html><body><div>stamp</div></body></html>", fragments[1])
    assert page.index("stamp") < page.index("<img")
    assert with_fragments(None, ["<i>x</i>"]).count("<i>x</i>") == 1


def test_image_must_be_png_data_uri() -> None:
    from kchs_engine.render.documents import Placement

    with pytest.raises(RenderContentError, match="PNG"):
        image_fragment(
            _image(dataUri="javascript:alert(1)"), Placement(0, 0.5, 0.5, "fallback"), A4
        )


def test_encrypted_pdf_is_content_error(tmp_path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.encrypt(user_password="secret", owner_password="owner")
    source = tmp_path / "locked.pdf"
    with source.open("wb") as stream:
        writer.write(stream)
    with pytest.raises(RenderContentError, match="паролем"):
        page_sizes(source, "first")


def test_broken_pdf_is_content_error(tmp_path: Path) -> None:
    source = tmp_path / "broken.pdf"
    source.write_bytes(b"%PDF-1.4\nnot really")
    with pytest.raises(RenderContentError):
        page_sizes(source, "first")


# ─── Задание целиком (api и хранилище — подделки) ────────────────────────────


class FakeApi:
    def __init__(self, start: dict[str, Any]) -> None:
        self.start = start
        self.done: list[dict[str, Any]] = []


@pytest.fixture
def fake_io(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> dict[str, Any]:
    state: dict[str, Any] = {"objects": {}, "uploaded": {}}

    async def object_size(bucket: str, key: str) -> int:
        return len(state["objects"][(bucket, key)])

    async def download(bucket: str, key: str, target: Path) -> Path:
        target.write_bytes(state["objects"][(bucket, key)])
        return target

    async def upload(bucket: str, key: str, source: Path, content_type: str) -> None:
        state["uploaded"][(bucket, key)] = (source.read_bytes(), content_type)

    monkeypatch.setattr(documents, "object_size", object_size)
    monkeypatch.setattr(documents, "download", download)
    monkeypatch.setattr(documents, "upload", upload)
    return state


def _wire_api(monkeypatch: pytest.MonkeyPatch, api: FakeApi) -> None:
    async def start(render_id: str) -> dict[str, Any]:
        return api.start

    async def done(render_id: str, payload: dict[str, Any]) -> dict[str, Any]:
        api.done.append(payload)
        return {"ok": True, "stale": False}

    monkeypatch.setattr(documents, "document_render_start", start)
    monkeypatch.setattr(documents, "document_render_done", done)


def test_job_fill_uploads_docx_and_reports(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, fake_io: dict[str, Any]
) -> None:
    template = _docx(tmp_path / "t.docx", ["{{ doc.subject }}"])
    fake_io["objects"][("files", "tpl")] = template.read_bytes()
    api = FakeApi(
        {
            "status": "render",
            "plan": {
                "kind": "docx",
                "template": {"bucket": "files", "storageKey": "tpl"},
                "context": {"doc": {"subject": "Ответ на запрос"}},
            },
            "target": {
                "bucket": "files",
                "storageKey": "out/render.docx",
                "fileName": "Ответ.docx",
                "contentType": "application/docx",
            },
        }
    )
    _wire_api(monkeypatch, api)
    result = asyncio.run(documents.document_render({"renderId": "r1"}))
    assert result["kind"] == "docx"
    data, content_type = fake_io["uploaded"][("files", "out/render.docx")]
    assert content_type == "application/docx"
    filled = tmp_path / "filled.docx"
    filled.write_bytes(data)
    assert "Ответ на запрос" in _text(filled)
    assert api.done == [{"status": "ready", "size": len(data), "pages": None}]


def test_job_inspect_reports_placeholders(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, fake_io: dict[str, Any]
) -> None:
    template = _docx(tmp_path / "t.docx", ["{{ doc.subject }} {{ author.position }}"])
    fake_io["objects"][("files", "tpl")] = template.read_bytes()
    api = FakeApi(
        {
            "status": "render",
            "plan": {"kind": "inspect", "template": {"bucket": "files", "storageKey": "tpl"}},
            "target": None,
        }
    )
    _wire_api(monkeypatch, api)
    asyncio.run(documents.document_render({"renderId": "r2"}))
    assert api.done == [{"status": "ready", "placeholders": ["author.position", "doc.subject"]}]


def test_job_bad_template_reports_failure_without_retry(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, fake_io: dict[str, Any]
) -> None:
    template = _docx(tmp_path / "t.docx", ["{% if %}"])
    fake_io["objects"][("files", "tpl")] = template.read_bytes()
    api = FakeApi(
        {
            "status": "render",
            "plan": {
                "kind": "docx",
                "template": {"bucket": "files", "storageKey": "tpl"},
                "context": {},
            },
            "target": {
                "bucket": "files",
                "storageKey": "out",
                "fileName": "x.docx",
                "contentType": "x",
            },
        }
    )
    _wire_api(monkeypatch, api)
    result = asyncio.run(documents.document_render({"renderId": "r3"}))
    assert "failed" in result
    assert api.done[0]["status"] == "failed"
    assert "Ошибка в шаблоне" in api.done[0]["error"]
    assert fake_io["uploaded"] == {}


def test_job_skip_does_nothing(monkeypatch: pytest.MonkeyPatch, fake_io: dict[str, Any]) -> None:
    api = FakeApi({"status": "skip", "reason": "no_access"})
    _wire_api(monkeypatch, api)
    assert asyncio.run(documents.document_render({"renderId": "r4"})) == {"skipped": "no_access"}
    assert api.done == []


# ─── Chromium (образ движка) ─────────────────────────────────────────────────

browser = pytest.mark.browser


def _pixel(pdf: Path, x_ratio: float, y_ratio: float, workdir: Path) -> tuple[int, int, int]:
    """Цвет точки первого листа: pdftoppm → PPM, доли ширины и высоты от левого верхнего угла."""
    prefix = workdir / "raster"
    subprocess.run(
        ["pdftoppm", "-r", "30", "-f", "1", "-l", "1", str(pdf), str(prefix)], check=True
    )
    ppm = next(workdir.glob("raster*.ppm")).read_bytes()
    parts = ppm.split(b"\n", 3)
    width, height = (int(value) for value in parts[1].split())
    pixels = parts[3]
    x = int(width * x_ratio)
    y = int(height * y_ratio)
    offset = (y * width + x) * 3
    return pixels[offset], pixels[offset + 1], pixels[offset + 2]


@browser
@pytest.mark.skipif(shutil.which("pdftoppm") is None, reason="нет poppler")
async def test_stamp_overlay_keeps_page_visible(tmp_path: Path, fake_io: dict[str, Any]) -> None:
    """Наложение прозрачно вне штампа: лист под ним виден, штамп — в правом нижнем углу."""
    from kchs_engine.render import report

    green = "<html><body style='margin:0;background:#00ff00'>&nbsp;</body></html>"
    stamp = (
        "<html><body style='margin:0;background:transparent'>"
        "<div style='position:absolute;right:14mm;bottom:12mm;width:50mm;height:20mm;"
        "background:#0000ff'></div></body></html>"
    )
    plan = {
        "kind": "overlay",
        "source": {
            "bucket": "files",
            "storageKey": "scan",
            "name": "scan.pdf",
            "mime": "application/pdf",
        },
        "html": stamp,
        "pages": "first",
    }
    # Один цикл событий на весь сценарий: браузер движка — одиночка процесса
    try:
        source_pdf, _ = await documents.render_html(
            {"html": green, "orientation": "portrait", "labels": {}}, tmp_path
        )
        fake_io["objects"][("files", "scan")] = source_pdf.read_bytes()
        result, pages, placements = await documents.render_overlay(plan, tmp_path)
    finally:
        await report.close_browser()
    assert pages >= 1
    assert placements is None
    top_left = _pixel(result, 0.3, 0.3, tmp_path)
    assert top_left[1] > 200 and top_left[2] < 60, top_left
    for leftover in tmp_path.glob("raster*.ppm"):
        leftover.unlink()
    corner = _pixel(result, 0.8, 0.92, tmp_path)
    assert corner[2] > 200 and corner[1] < 60, corner


@browser
async def test_html_form_blocks_network(tmp_path: Path) -> None:
    from kchs_engine.render import report

    page = "<html><body><h1>Карточка</h1><img src='http://127.0.0.1:9/leak.png'></body></html>"
    try:
        path, pages = await documents.render_html(
            {
                "html": page,
                "orientation": "portrait",
                "footer": "kchs",
                "labels": {"page": "Лист", "of": "из"},
            },
            tmp_path,
        )
    finally:
        await report.close_browser()
    assert path.read_bytes().startswith(b"%PDF")
    assert pages == 1


@browser
@pytest.mark.skipif(shutil.which("pdftoppm") is None, reason="нет poppler")
async def test_facsimile_lands_left_of_cyrillic_name(
    tmp_path: Path, fake_io: dict[str, Any]
) -> None:
    """Факсимиле встаёт левее фамилии, напечатанной кириллицей, и видно на листе (ADR-0198)."""
    import base64

    from PIL import Image

    from kchs_engine.render import report

    picture = io.BytesIO()
    Image.new("RGBA", (300, 120), (255, 0, 0, 255)).save(picture, format="PNG")
    letter = (
        "<html><body style='margin:25mm;font:14pt serif'>"
        "<p>Компания благодарна Вам за внимание к нашим продуктам.</p>"
        "<p style='text-align:right;margin-top:20mm'>Саидзода Т. Н.</p></body></html>"
    )
    stamp = (
        "<html><body style='margin:0;background:transparent'>"
        "<div style='position:absolute;left:12mm;bottom:10mm;width:60mm;height:18mm;"
        "background:#0000ff'></div></body></html>"
    )
    plan = {
        "kind": "overlay",
        "source": {
            "bucket": "files",
            "storageKey": "letter",
            "name": "letter.pdf",
            "mime": "application/pdf",
        },
        "html": stamp,
        "pages": "last",
        "images": [
            {
                "id": "sig-1",
                "dataUri": "data:image/png;base64," + base64.b64encode(picture.getvalue()).decode(),
                "widthMm": 40.0,
                "heightMm": 16.0,
                "anchors": [
                    {"text": "[[SIGN]]", "mode": "cover"},
                    {"text": "Саидзода Т. Н.", "mode": "left"},
                ],
                "manual": None,
                "fallback": {"x": 0.7, "y": 0.75},
            }
        ],
    }
    try:
        source_pdf, _ = await documents.render_html(
            {"html": letter, "orientation": "portrait", "labels": {}}, tmp_path
        )
        fake_io["objects"][("files", "letter")] = source_pdf.read_bytes()
        result, pages, placements = await documents.render_overlay(plan, tmp_path)
    finally:
        await report.close_browser()
    assert pages >= 1
    assert placements is not None and placements[0]["source"] == "name"
    spot = placements[0]
    red = _pixel(result, spot["x"], spot["y"], tmp_path)
    assert red[0] > 200 and red[1] < 80 and red[2] < 80, red
    for leftover in tmp_path.glob("raster*.ppm"):
        leftover.unlink()
    # Штамп — тоже на последнем листе
    corner = _pixel(result, 0.15, 0.95, tmp_path)
    assert corner[2] > 200 and corner[1] < 60, corner
