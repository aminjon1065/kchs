"""Сводка «10 сола», классификатор, итоговый реестр и файл сверки для владельца.

Выход: /out/incident_types.csv, /out/stats_10y.csv, /out/incidents.csv (Disaster + события
2021–2026 из актов), /out/Сверка.xlsx.
"""

from __future__ import annotations

import collections
import csv
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, "/tools")
from common import KINDS, REPORT_CATEGORIES, load_territories  # noqa: E402

OUT = "/out/"
CATEGORY_COLUMNS = ["EQ", "MF", "AV", "WIND", "WATER", "SLIDE", "SNOW", "RAIN", "DROUGHT", "HYDRO", "WET", "BIO"]


def read(name: str) -> list[dict]:
    with open(OUT + name, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def write(name: str, rows: list[dict]) -> None:
    with open(OUT + name, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def official_stats() -> tuple[list[dict], list[str]]:
    """«10 сола»: на каждый год три строки — количество ЧС, погибшие, ущерб (сомони)."""
    wb = openpyxl.load_workbook("/data/10 сола 2013-2024.xlsx", data_only=True)
    result: dict[tuple, dict] = {}
    notes = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        year = None
        for r in rows[4:]:
            first = str(r[0] or "").strip()
            if first.isdigit():
                year = int(first)
            elif first:
                year = None  # «Ҳамагӣ» — итог за все годы, а не год
            label = str(r[1] or "").strip().lower()
            if not year or not label:
                continue
            measure = "count" if label.startswith("миқдор") else "deaths" if label.startswith("фавт") else "damage" if label.startswith("маблағ") else None
            if not measure:
                continue
            for i, cat in enumerate(CATEGORY_COLUMNS + ["TOTAL"]):
                value = r[2 + i]
                if value in (None, ""):
                    continue
                key = (year, cat)
                item = result.setdefault(key, {"year": year, "category": cat, "count": "", "deaths": "", "damage": ""})
                if item[measure] not in ("", value) and ws.title != wb.worksheets[0].title:
                    notes.append(f"{year} {cat} {measure}: лист «{ws.title}» {value} ≠ {item[measure]}")
                    continue
                if item[measure] == "":
                    item[measure] = value
    return sorted(result.values(), key=lambda x: (x["year"], x["category"])), notes


def main() -> None:
    territories = load_territories()
    names = {c: (t.ru or c) for c, t in territories.items()}
    report = {code: (ru, tg) for code, ru, tg in REPORT_CATEGORIES}

    # ── Классификатор ──────────────────────────────────────────────────────
    types = [{"code": code, "name": ru, "name_tg": tg, "group_name": group, "report_category": cat,
              "report_category_name": report[cat][0]} for code, ru, tg, group, cat in KINDS]
    write("incident_types.csv", types)
    kind_category = {t["code"]: t["report_category"] for t in types}
    kind_name = {t["code"]: t["name"] for t in types}

    # ── Реестр: Disaster + события 2021–2026 из актов ──────────────────────
    disaster = read("incidents_disaster.csv")
    from_acts = read("incidents_damage.csv")
    incidents = disaster + from_acts
    write("incidents.csv", incidents)

    # ── Официальная статистика ─────────────────────────────────────────────
    stats, stat_notes = official_stats()
    write("stats_10y.csv", [{**s, "category_name": report.get(s["category"], ("Всего",))[0]} for s in stats])

    acts = read("damage_assessments.csv")
    reg_count = collections.Counter()
    reg_deaths = collections.Counter()
    for r in incidents:
        year = int(r["occurred_at"][:4])
        cat = kind_category[r["type_code"]]
        reg_count[(year, cat)] += int(r.get("occurrences") or 1)
        reg_deaths[(year, cat)] += int(r["deaths"] or 0)
    act_damage = collections.Counter()
    for a in acts:
        act_damage[int(a["year"])] += float(a["damage"] or 0)

    # ── Сверка.xlsx ─────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE7F3")
    warn_fill = PatternFill("solid", fgColor="FCE8C8")

    def sheet(title: str, header: list[str], rows: list[list], widths: list[int], warn=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in rows:
            ws.append(row)
            if warn and warn(row):
                for cell in ws[ws.max_row]:
                    cell.fill = warn_fill
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    # Итог
    ws = wb.active
    ws.title = "Итог"
    by_level = collections.Counter(territories[r["territory"]].level for r in incidents)
    lines = [
        ("Что будет загружено", ""),
        ("Реестр «Происшествия» (пакет ЧС)", f"{len(incidents)} событий: {len(disaster)} из таблиц Disaster (1988–2020) и {len(from_acts)} из актов ущерба (2021–2026)"),
        ("Привязка к территориям", f"к району — {by_level['district']}, к области — {by_level['region']}, ко всей республике — {by_level['country']} (место не указано, вся РТ или за рубежом)"),
        ("Точность дат", ", ".join(f"{k}: {v}" for k, v in collections.Counter(r['date_precision'] for r in incidents).items()) + " (day — день, month — месяц, year — только год)"),
        ("Погибшие", f"{sum(int(r['deaths'] or 0) for r in incidents)} человек; у {sum(1 for r in incidents if r['deaths_from_text'])} событий 1992–2016 и актов число взято из описания"),
        ("Классификатор «Типы происшествий»", f"{len(types)} видов: 22 вида пакета ЧС и новые; у каждого — категория сводки «10 сола»"),
        ("Реестр «Оценка ущерба»", f"{len(acts)} актов 2013–2026; привязано к событиям реестра: {sum(1 for a in acts if a['incident_code'])}"),
        ("Статистика «10 сола»", f"{len(stats)} строк (год × категория): количество, погибшие, ущерб — официальные цифры"),
        ("Слои карт", "13 слоёв геобазы ArcGIS (лист «Слои карт»); рельеф — после доработки подложки (см. ниже)"),
        ("", ""),
        ("Что проверить", ""),
        ("Лист «Виды ЧС»", "как сведены написания вида к классификатору; поправьте код в колонке «код», если не согласны"),
        ("Лист «Районы»", "оранжевые строки — район не распознан, несколько районов или район из другой области"),
        ("Лист «Сверка с 10 сола»", "сколько событий насчитал реестр и сколько в официальной сводке; расхождения — повод проверить"),
        ("Листы «Акты без суммы», «Без года»", "строки, которые загрузятся неполными или не загрузятся"),
        ("", ""),
        ("Решения при разборе", ""),
        ("Дубли файлов", "основная — самая полная версия года; из копий взяты только строки, которых нет в основной (лист «Дубли»)"),
        ("Ущерб до 2001 года", "в рублях разного масштаба — только текстом «как в источнике»; числом (сомони) — с 2001 года"),
        ("Старые буквы", "њ ї ѓ ќ љ ў заменены на ҳ ӣ ғ қ ҷ ӯ — иначе не работают поиск и привязка к районам"),
        ("Переименованные районы", "Джиргатал → Лахш, Ганчи → Деваштич, Колхозобод → Дж. Балхи и т. д.; «Бохтар» до 2018 года — Бохтарский район (ныне Кушониён)"),
        ("Рельеф", "платформа пока берёт растровые подложки только по внешнему адресу; для отмывки рельефа нужна доработка — растровая подложка из файла"),
        ("База ESTJ (2010)", "классификатор и бланки донесений — образец; записей о событиях в ней почти нет"),
    ]
    for a, b in lines:
        ws.append([a, b])
        if a and not b:
            ws[ws.max_row][0].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 120

    # Классификатор
    sheet("Классификатор", ["код", "вид", "на таджикском", "группа", "категория сводки «10 сола»"],
          [[t["code"], t["name"], t["name_tg"], t["group_name"], t["report_category_name"]] for t in types],
          [14, 48, 44, 20, 32])

    # Виды ЧС: написание → код
    spell = collections.Counter()
    for r in incidents:
        spell[(r["type_raw"].lower(), r["type_code"])] += 1
    sheet("Виды ЧС", ["вид в источнике", "код", "вид в классификаторе", "событий"],
          [[raw, code, kind_name[code], n] for (raw, code), n in spell.most_common()],
          [60, 14, 48, 10], warn=lambda row: row[1] == "OTHER")

    # Районы
    places = collections.Counter()
    for r in incidents:
        places[(r["region_raw"], r["district_raw"][:80], r["territory"], r["territory_note"])] += 1
    sheet("Районы", ["область в источнике", "район в источнике", "территория", "код", "замечание", "событий"],
          [[p, d, names.get(c, ""), c, note, n] for (p, d, c, note), n in places.most_common()],
          [24, 50, 34, 12, 46, 10], warn=lambda row: bool(row[4]))

    # Сверка с 10 сола
    official = {(s["year"], s["category"]): s for s in stats}
    rows = []
    for year in range(2013, 2025):
        for cat in CATEGORY_COLUMNS:
            o = official.get((year, cat), {})
            oc, od = o.get("count", ""), o.get("deaths", "")
            rc, rd = reg_count.get((year, cat), 0), reg_deaths.get((year, cat), 0)
            if oc in ("", 0) and rc == 0:
                continue
            rows.append([year, report[cat][0], oc, rc if year <= 2020 or rc else "", od, rd])
        total = official.get((year, "TOTAL"), {})
        rows.append([year, "ВСЕГО (ущерб, тыс. сомони: сводка / акты)",
                     total.get("count", ""), sum(reg_count.get((year, c), 0) for c in CATEGORY_COLUMNS),
                     round(float(total.get("damage") or 0) / 1000, 1), round(act_damage.get(year, 0) / 1000, 1)])
    sheet("Сверка с 10 сола", ["год", "категория", "ЧС по сводке", "ЧС в реестре (с учётом «N тарма»)", "погибших по сводке", "погибших в реестре"],
          rows, [8, 44, 14, 14, 18, 18],
          warn=lambda row: isinstance(row[2], (int, float)) and isinstance(row[3], int) and row[2] and abs(row[3] - row[2]) > max(2, 0.3 * row[2]))
    if stat_notes:
        ws = wb["Сверка с 10 сола"]
        ws.append([])
        ws.append(["Листы сводки расходятся:"])
        for n in stat_notes[:40]:
            ws.append([n])

    # Дубли
    dedupe = list(csv.reader(open(OUT + "review_dedupe.csv", encoding="utf-8")))
    sheet("Дубли", dedupe[0], dedupe[1:], [64, 18, 20, 16, 16])

    # Акты без суммы или даты
    sheet("Акты без суммы", ["акт", "год", "район", "дата", "погибли", "текст"],
          [[a["code"], a["year"], names.get(a["territory"], a["territory"]), a["event_date"], a["deaths"], a["text"][:400]]
           for a in acts if not a["damage"] or not a["event_date"]],
          [12, 8, 26, 12, 9, 140])

    # Без года
    no_year = list(csv.reader(open(OUT + "review_no_year.csv", encoding="utf-8")))
    sheet("Без года", no_year[0], no_year[1:], [40, 10, 10, 20, 20, 120])

    # Слои карт
    gis = [r for r in csv.reader(open(OUT + "review_gis.csv", encoding="utf-8")) if r]
    sheet("Слои карт", gis[0], gis[1:], [26, 28, 30, 10, 18, 60, 30])

    wb.save(OUT + "Сверка.xlsx")
    print("incidents.csv:", len(incidents), "| types:", len(types), "| stats:", len(stats), "| расхождения листов:", len(stat_notes))


if __name__ == "__main__":
    main()
